"""Preenchimento da aba "Remessas Porto" a partir da expedicao.

Reescreve a area de dados inteira, com os subtotais de cada dia, e devolve
as chaves de acesso e os valores unitarios ja digitados para as suas notas.
Pode rodar todo dia: ela reconcilia, nao acrescenta.

VL. UNITARIO: o valor do config.ini e so o padrao para nota NOVA. Depois de
lancada, o que estiver na coluna G daquela nota e mantido - pode digitar o
preco certo de cada carregamento direto na planilha.
"""

from collections import defaultdict

import openpyxl

from . import expedicao
from .planilha import (
    ErroDeUso,
    achar_rodape,
    aplicar_formato,
    coluna,
    fazer_backup,
    garantir_brl,
    gravar,
    guardar_formato,
    inserir_linhas,
    limpar_conteudo,
    como_numero,
    texto,
)

ABA = "Remessas Porto"
LIN_INI = 4

C_DATA, C_NF, C_PLACA, C_TRANSP = 1, 2, 3, 4
C_LIQ, C_BRUTO, C_UNIT, C_TOTAL = 5, 6, 7, 8
C_CFOP, C_OPER, C_VLDIA, C_KGDIA = 9, 10, 11, 12
C_CHAVE, C_LACRE = 13, 14
C_ULT = 14


def atualizar(cfg, relato=print):
    caminho_exp = cfg.expedicao
    caminho_rem = cfg.remessa

    relato(f"Lendo a expedicao: {caminho_exp.name}")
    wb_exp = openpyxl.load_workbook(caminho_exp, data_only=True, read_only=False)
    notas, sem_peso = expedicao.ler(wb_exp, cfg.ano)
    wb_exp.close()
    relato(f"  {len(notas)} nota(s) em {len(set(n.data for n in notas.values()))} dia(s)")

    relato(f"Abrindo a remessa: {caminho_rem.name}")
    wb = openpyxl.load_workbook(caminho_rem)
    if ABA not in wb.sheetnames:
        raise ErroDeUso(
            f'Nao encontrei a aba "{ABA}" em {caminho_rem.name}.\n'
            "Confira se o caminho aponta para a planilha de remessa."
        )
    ws = wb[ABA]

    rodape = achar_rodape(ws, coluna=C_TRANSP)
    if rodape == 0:
        raise ErroDeUso(
            'Nao encontrei a linha de rodape ("VOLUME EMBARCADO") na coluna D.\n'
            "A rotina precisa dela para saber onde termina a area de dados.\n"
            "Nada foi alterado."
        )
    lin_fim = rodape - 1

    chaves = _guardar_chaves(ws, LIN_INI, lin_fim)
    relato(f"  {len(chaves)} chave(s) de acesso preservada(s)")
    unitarios = _guardar_unitarios(ws, LIN_INI, lin_fim)
    relato(f"  {len(unitarios)} valor(es) unitario(s) preservado(s)")

    # linhas-modelo de formatacao, fotografadas antes de mexer em nada
    fmt_dado = guardar_formato(ws, LIN_INI, C_ULT)
    lin_sub = _achar_primeiro_subtotal(ws, LIN_INI, lin_fim)
    fmt_sub = guardar_formato(ws, lin_sub or LIN_INI, C_ULT)

    dias = sorted({n.data for n in notas.values()})
    precisa = len(notas) + len(dias)
    disponivel = lin_fim - LIN_INI + 1

    if precisa > disponivel:
        faltam = precisa - disponivel
        # inserir DENTRO da area (na ultima linha dela) para que as formulas do
        # rodape, que somam de LIN_INI ate lin_fim, estiquem junto
        relato(f"  inserindo {faltam} linha(s) na area de dados")
        inserir_linhas(wb, ws, lin_fim, faltam)
        lin_fim += faltam

    backup = fazer_backup(caminho_rem)
    relato(f"  copia de seguranca: Backup\\{backup.name}")

    limpar_conteudo(ws, LIN_INI, lin_fim, C_ULT)

    por_dia = defaultdict(list)
    for nf, nota in notas.items():
        por_dia[nota.data].append(nf)

    unit_padrao = []          # notas novas, que receberam o valor do config.ini
    lin = LIN_INI
    for data in dias:
        inicio_bloco = lin
        for nf in sorted(por_dia[data]):
            nota = notas[nf]
            aplicar_formato(ws, lin, fmt_dado)
            transp = nota.transportadora or cfg.transportadora_padrao
            ws.cell(lin, C_DATA).value = data
            ws.cell(lin, C_NF).value = nf
            ws.cell(lin, C_PLACA).value = nota.placa
            ws.cell(lin, C_TRANSP).value = transp
            ws.cell(lin, C_LIQ).value = nota.liquido
            ws.cell(lin, C_BRUTO).value = nota.liquido
            if nf in unitarios:
                ws.cell(lin, C_UNIT).value = unitarios[nf]
            else:
                ws.cell(lin, C_UNIT).value = cfg.valor_unitario
                unit_padrao.append(nf)
            ws.cell(lin, C_TOTAL).value = f"=G{lin}*E{lin}"
            garantir_brl(ws.cell(lin, C_UNIT))
            garantir_brl(ws.cell(lin, C_TOTAL))
            ws.cell(lin, C_CFOP).value = cfg.cfop_remessa
            ws.cell(lin, C_OPER).value = cfg.operacao_remessa
            ws.cell(lin, C_LACRE).value = nota.lacre
            if nf in chaves:
                ws.cell(lin, C_CHAVE).value = chaves[nf]
            lin += 1

        aplicar_formato(ws, lin, fmt_sub)
        ws.cell(lin, C_VLDIA).value = f"=SUM(H{inicio_bloco}:H{lin - 1})"
        garantir_brl(ws.cell(lin, C_VLDIA))
        ws.cell(lin, C_KGDIA).value = f"=SUM(E{inicio_bloco}:E{lin - 1})"
        _destacar_subtotal(ws, lin)
        lin += 1

    # sobras: linha-modelo em branco, pronta para o proximo dia
    while lin <= lin_fim:
        aplicar_formato(ws, lin, fmt_dado)
        ws.cell(lin, C_TRANSP).value = cfg.transportadora_padrao
        ws.cell(lin, C_UNIT).value = cfg.valor_unitario
        ws.cell(lin, C_TOTAL).value = f"=G{lin}*E{lin}"
        ws.cell(lin, C_CFOP).value = cfg.cfop_remessa
        ws.cell(lin, C_OPER).value = cfg.operacao_remessa
        lin += 1

    gravar(wb, caminho_rem)

    com_chave = sum(1 for nf in notas if nf in chaves)
    return {
        "notas": len(notas),
        "dias": len(dias),
        "com_chave": com_chave,
        "sem_chave": len(notas) - com_chave,
        "sem_peso": sem_peso,
        "unit_padrao": unit_padrao,
        "valor_padrao": cfg.valor_unitario,
        "backup": backup,
        "linhas_inseridas": max(0, precisa - disponivel),
    }


def _guardar_chaves(ws, ini, fim):
    chaves = {}
    for r in range(ini, fim + 1):
        nf = como_numero(ws.cell(r, C_NF).value)
        if nf is None:
            continue
        chave = texto(ws.cell(r, C_CHAVE).value)
        if chave and nf not in chaves:
            chaves[nf] = chave
    return chaves


def _guardar_unitarios(ws, ini, fim):
    """{nf: valor unitario} das notas que ja estao na aba.

    Guarda o que estiver digitado na coluna VL. UNITARIO de cada nota, para
    que rodar de novo nao volte tudo para o valor padrao do config.ini.
    """
    unitarios = {}
    for r in range(ini, fim + 1):
        nf = como_numero(ws.cell(r, C_NF).value)
        if nf is None:
            continue
        unit = como_numero(ws.cell(r, C_UNIT).value)
        if unit is not None and nf not in unitarios:
            unitarios[nf] = unit
    return unitarios


def _achar_primeiro_subtotal(ws, ini, fim):
    """Linha de subtotal = sem Nº NF e com formula na coluna VL. TOTAL DO DIA."""
    for r in range(ini, fim + 1):
        if texto(ws.cell(r, C_NF).value):
            continue
        v = ws.cell(r, C_VLDIA).value
        if isinstance(v, str) and v.startswith("="):
            return r
    return 0


def _destacar_subtotal(ws, lin):
    """Fecha o bloco do dia.

    Borda inferior espessa atravessando a linha inteira, de A ate N, e
    negrito nas duas colunas de total (VL. TOTAL DO DIA e KG. TOTAL DIA).
    """
    from copy import copy

    from openpyxl.styles import Side

    for c in range(1, C_ULT + 1):
        celula = ws.cell(lin, c)
        borda = copy(celula.border)
        borda.bottom = Side(style="thick")
        celula.border = borda
        if c in (C_VLDIA, C_KGDIA):
            fonte = copy(celula.font)
            fonte.bold = True
            celula.font = fonte
