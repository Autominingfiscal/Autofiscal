"""Preenchimento da coluna CHAVE ACESSO a partir dos PDFs da rede.

Percorre as linhas da "Remessas Porto" que estao sem chave, procura na pasta
o PDF daquela nota, le a chave de dentro do arquivo e grava no formato de
sempre. Linha que ja tem chave nao e tocada.
"""

import re

import openpyxl

from . import pdf
from .planilha import (
    ErroDeUso,
    achar_rodape,
    como_numero,
    fazer_backup,
    gravar,
    texto,
)
from .remessas import ABA, C_CHAVE, C_NF, C_TRANSP, LIN_INI

_NUM_NO_NOME = re.compile(r"\d{4,}")


def indexar_pdfs(pasta, incluir_subpastas=True):
    """{numero: [caminhos]} por cada numero de 4+ digitos do nome."""
    indice = {}
    pastas = 0
    arquivos = 0
    caminhos = pasta.rglob("*") if incluir_subpastas else pasta.glob("*")

    vistas = set()
    for item in caminhos:
        try:
            if item.is_dir():
                if item not in vistas:
                    vistas.add(item)
                    pastas += 1
                continue
            if item.suffix.lower() != ".pdf":
                continue
        except OSError:
            continue  # pasta sem permissao de leitura: pula em silencio
        # TODOS os numeros do nome, nao so o primeiro: em "2026-09 NF 19007.pdf"
        # o primeiro e o ano. Indexar a mais e seguro, porque a chave lida do PDF
        # so e aceita se a nota dentro dela for a esperada.
        numeros = {float(n) for n in _NUM_NO_NOME.findall(item.stem)}
        if not numeros:
            continue
        for n in numeros:
            indice.setdefault(n, []).append(item)
        arquivos += 1

    return indice, pastas + 1, arquivos


def preencher(cfg, relato=print):
    caminho_rem = cfg.remessa
    pasta = cfg.pdfs

    relato(f"Varrendo {pasta}")
    indice, n_pastas, n_arquivos = indexar_pdfs(pasta, cfg.incluir_subpastas)
    if not indice:
        extra = " (nem nas subpastas)" if cfg.incluir_subpastas else ""
        raise ErroDeUso(
            f"Nao achei nenhum PDF{extra} em:\n  {pasta}\n\n"
            f"Pastas varridas: {n_pastas}"
        )
    relato(f"  {n_arquivos} PDF(s) indexado(s) em {n_pastas} pasta(s)")

    wb = openpyxl.load_workbook(caminho_rem)
    if ABA not in wb.sheetnames:
        raise ErroDeUso(f'Nao encontrei a aba "{ABA}" em {caminho_rem.name}.')
    ws = wb[ABA]

    rodape = achar_rodape(ws, coluna=C_TRANSP)
    if rodape <= LIN_INI:
        raise ErroDeUso(
            'Nao encontrei a linha de rodape ("VOLUME EMBARCADO") na coluna D.'
        )

    backup = fazer_backup(caminho_rem)
    relato(f"  copia de seguranca: Backup\\{backup.name}")

    preenchidas = ja_tinha = sem_pdf = nao_validaram = 0
    pendencias = []

    for r in range(LIN_INI, rodape):
        nf = como_numero(ws.cell(r, C_NF).value)
        if nf is None:
            continue

        if texto(ws.cell(r, C_CHAVE).value):
            ja_tinha += 1
            continue

        arquivos = indice.get(nf)
        if not arquivos:
            sem_pdf += 1
            pendencias.append(f"NF {int(nf)}: PDF nao encontrado na pasta")
            continue

        # o mesmo numero pode aparecer em mais de uma subpasta: tenta ate validar
        chave = ""
        primeiro_erro = ""
        for caminho in arquivos:
            chave, motivo = pdf.chave_do_pdf(caminho, nf)
            if chave:
                break
            if not primeiro_erro:
                primeiro_erro = motivo

        if chave:
            ws.cell(r, C_CHAVE).value = pdf.formatar(chave)
            preenchidas += 1
        else:
            nao_validaram += 1
            extra = f" ({len(arquivos)} arquivos com esse numero)" if len(arquivos) > 1 else ""
            pendencias.append(f"NF {int(nf)}: {primeiro_erro}{extra}")

    gravar(wb, caminho_rem)

    return {
        "preenchidas": preenchidas,
        "ja_tinha": ja_tinha,
        "sem_pdf": sem_pdf,
        "nao_validaram": nao_validaram,
        "pdfs": n_arquivos,
        "pastas": n_pastas,
        "subpastas": cfg.incluir_subpastas,
        "pendencias": pendencias,
        "backup": backup,
    }


def conferir_um(caminho, nf_esperada=None):
    """Teste de um PDF isolado, como a macro ConferirUmPDF."""
    chave, motivo = pdf.chave_do_pdf(caminho, nf_esperada)
    if chave:
        return {
            "chave": pdf.formatar(chave),
            "nota_na_chave": pdf.nota_da_chave(chave),
            "ok": True,
            "motivo": "",
        }
    return {"chave": "", "nota_na_chave": None, "ok": False, "motivo": motivo}
