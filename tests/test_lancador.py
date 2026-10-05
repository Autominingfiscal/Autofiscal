"""Lancador: leitura do ticket e da nota, e as tres etapas numa planilha de teste."""

import contextlib
import datetime as dt
import io
import os
import unittest

import _apoio

import openpyxl
from openpyxl.worksheet.table import Table

import etapa3_lancador as lan

CFG = {
    "alvos": {"TRUCADO": 48500, "LS 4 EIXOS": 58500},
    "turno2": dt.time(12, 0), "tol_exc": 80, "tol_sub": 1000,
    "produto": "", "salvar": False, "bipe": False,
}

EXPED = ["Nº", "Nº_NOTA", "TICKET DE PESAGEM", "CÓDIGO LACRE", "TRANSPORTADORA", "MOTORISTA",
         "PLACA DO CAVALO", "UF CAVALO", "PLACA DO REBOQUE", "UF REBOQUE", "MODELO",
         "PESO_ENTRADA", "PESO_SAÍDA", "HORA_ENTRADA", "HORA_SAÍDA", "HORA EMISSÃO NF"]
TURNO = ["CÓDIGO LACRE", "TRANSPORTADORA", "MOTORISTA", "PLACA DO CAVALO", "UF CAVALO",
         "PLACA DO REBOQUE", "UF REBOQUE", "MODELO", "PESO A CARREGAR"]


def ticket_pdf(caminho, campos, motorista):
    """PDF no jeito do FortesReport: rotulo (F1) e valor (F2) na mesma linha."""
    partes, y = [], 700
    for rotulo, valor in campos.items():
        partes.append(b"/F1 8 Tf 20 %d Td (%s) Tj" % (y, rotulo.encode("cp1252")))
        partes.append(b"/F2 8 Tf 90 %d Td (%s) Tj" % (y, valor.encode("cp1252")))
        y -= 12
    partes.append(b"/F2 8 Tf 60 108 Td (%s) Tj" % motorista.encode("cp1252"))
    partes.append(b"/F1 8 Tf 60 100 Td (Assinatura do Motorista) Tj")
    with open(caminho, "wb") as f:
        f.write(b"%PDF-1.3\nBT\n" + b"\n".join(partes) + b"\nET\n%%EOF\n")


class TestUtilidades(unittest.TestCase):
    def test_desescapar(self):
        self.assertEqual(lan._desescapar(rb"a\101\(b\)\9"), "aA(b)9")

    def test_data_hora(self):
        self.assertEqual(lan._data_hora("18/09/26 07:38"), dt.datetime(2026, 9, 18, 7, 38))
        self.assertIsNone(lan._data_hora("31/02/2026 10:00"))
        self.assertIsNone(lan._data_hora("sem data"))

    def test_peso_alvo(self):
        self.assertEqual(lan.peso_alvo(CFG, "LS 4 EIXOS"), 58500)
        self.assertEqual(lan.peso_alvo(CFG, "ls-4 eixos"), 58500)
        self.assertIsNone(lan.peso_alvo(CFG, "BITREM"))

    def test_hora_gravada_e_lida_pelo_dashboard(self):
        import painel_local
        momento = dt.datetime(2026, 9, 18, 7, 38, 16)
        self.assertAlmostEqual(painel_local.minutos(lan.hora_excel_s(momento)),
                               7 * 60 + 38 + 16 / 60)


class TestLerTicket(unittest.TestCase):
    def test_ticket_de_tara(self):
        caminho = os.path.join(_apoio.pasta_temporaria(self), "t.pdf")
        ticket_pdf(caminho, {
            "NºTicket": "19569", "Produto": "CONCENTRADO", "Transportadora": "RODOGRANEL",
            "Cavalo": "ABC1D23", "1º Reboque": "XYZ9876", "Data Entrada": "18/09/2026 07:04",
            "Líquido": "17.250", "Tara": "", "Bruto": "",
        }, "JOAO DA SILVA")
        t = lan.ler_ticket(caminho)
        self.assertEqual((t["numero"], t["tipo"], t["peso_entrada"]), (19569, "tara", 17250.0))
        self.assertEqual(t["motorista"], "JOAO DA SILVA")
        self.assertEqual(t["entrada"], dt.datetime(2026, 9, 18, 7, 4))

    def test_pdf_sem_ticket_devolve_none(self):
        caminho = os.path.join(_apoio.pasta_temporaria(self), "t.pdf")
        ticket_pdf(caminho, {"Produto": "CONCENTRADO"}, "")
        self.assertIsNone(lan.ler_ticket(caminho))


class TestLerNota(unittest.TestCase):
    def texto(self, chave):
        return ("DANFE CHAVE DE ACESSO %s PROTOCOLO DE AUTORIZACAO DE USO "
                "126260001234567 - 18/09/2026 09:15:30 TICKET 19569 PLACA ABC1D23 "
                "PESO LIQUIDO 41.230,000" % " ".join(chave[i:i + 4] for i in range(0, 44, 4)))

    def test_le_numero_protocolo_ticket_placa(self):
        n = lan.ler_nota("x", self.texto(_apoio.chave_valida(18999)))
        self.assertEqual(n["numero"], 18999)
        self.assertEqual(n["protocolo"], dt.datetime(2026, 9, 18, 9, 15, 30))
        self.assertIn(19569, n["tickets"])
        self.assertIn("ABC1D23", n["placas"])
        self.assertIn(41230.0, n["pesos"])

    def test_chave_com_dv_errado_e_recusada(self):
        chave = _apoio.chave_valida(18999)
        errada = chave[:-1] + str((int(chave[-1]) + 1) % 10)
        self.assertIsNone(lan.ler_nota("x", self.texto(errada)))


class TestEtapasNaPlanilha(unittest.TestCase):
    """Tara -> completo -> nota, no modo --teste (planilha .xlsx, sem Excel)."""

    def setUp(self):
        wb = openpyxl.Workbook()
        ws = wb.active
        ws.title = "18.09"
        for c, cab in enumerate(EXPED, start=1):
            ws.cell(1, c).value = cab
        ws.add_table(Table(displayName="Expedicao", ref="A1:P4"))
        for linha_cab, nome in ((10, "Matutino"), (20, "Vespertino")):
            for c, cab in enumerate(TURNO, start=1):
                ws.cell(linha_cab, c).value = cab
            ws.add_table(Table(displayName=nome, ref="A%d:I%d" % (linha_cab, linha_cab + 1)))
        ws.cell(11, 2).value = "RODOGRANEL"
        ws.cell(11, 3).value = "JOAO DA SILVA"
        ws.cell(11, 4).value = "ABC-1D23"
        ws.cell(11, 8).value = "LS 4 EIXOS"
        caminho = os.path.join(_apoio.pasta_temporaria(self), "EXPED.xlsx")
        wb.save(caminho)
        self.planilha = lan.PlanilhaTeste(caminho)
        self.ws = self.planilha.wb["18.09"]

    def quieto(self, funcao, *args):
        with contextlib.redirect_stdout(io.StringIO()):
            return funcao(*args)

    def ticket(self, **extra):
        t = {"numero": 19569, "cavalo": "ABC1D23", "reboque": "", "motorista": "JOAO DA SILVA",
             "transportadora": "RODOGRANEL", "entrada": dt.datetime(2026, 9, 18, 7, 4),
             "saida": None, "peso_entrada": 17250.0, "tara": 0.0, "bruto": 0.0, "tipo": "tara"}
        t.update(extra)
        return t

    def col(self, nome):
        return EXPED.index(nome) + 1

    def test_tres_etapas(self):
        r = self.quieto(lan.lancar_tara, self.planilha, CFG, self.ticket())
        self.assertEqual(r[:3], ("LANÇADO", "18.09", 2))
        self.assertEqual(self.ws.cell(2, self.col("TICKET DE PESAGEM")).value, 19569)
        self.assertEqual(self.ws.cell(2, self.col("PESO_ENTRADA")).value, 17250)
        self.assertEqual(self.ws.cell(11, 9).value, 58500 - 17250)     # PESO A CARREGAR

        # a mesma tara de novo nao duplica
        self.planilha.limpar_cache()
        r = self.quieto(lan.lancar_tara, self.planilha, CFG, self.ticket())
        self.assertEqual(r[0], "JÁ LANÇADO")

        completo = self.ticket(tipo="completo", tara=17250.0, bruto=58700.0,
                               saida=dt.datetime(2026, 9, 18, 8, 30))
        r = self.quieto(lan.lancar_completo, self.planilha, CFG, completo)
        self.assertEqual(r[0], "SAÍDA - ALERTA")                      # 200 kg acima: vermelho
        celula = self.ws.cell(2, self.col("PESO_SAÍDA"))
        self.assertEqual(celula.value, 58700)
        self.assertEqual(celula.fill.fgColor.rgb[-6:], "FF0000")

        nota = lan.ler_nota("x", TestLerNota.texto(None, _apoio.chave_valida(18999)))
        r = self.quieto(lan.lancar_nota, self.planilha, CFG, nota)
        self.assertEqual(r[0], "LANÇADA")
        self.assertEqual(self.ws.cell(2, self.col("Nº_NOTA")).value, 18999)
        self.assertAlmostEqual(self.ws.cell(2, self.col("HORA EMISSÃO NF")).value,
                               (9 * 3600 + 15 * 60 + 30) / 86400)

    def test_caminhao_fora_do_turno_espera(self):
        r = self.quieto(lan.lancar_tara, self.planilha, CFG, self.ticket(cavalo="ZZZ9Z99"))
        self.assertEqual(r[0], "ESPERANDO")


class TestObservacoesAdicionais(unittest.TestCase):
    """Coluna X: 'TICKET N°: / PLACA CAVALO: / PLACA REBOQUE: / LACRES N°:' da propria linha."""

    CAB = EXPED + ["OBSERVAÇÕES\nADICIONAIS"]      # com a quebra de linha, como na planilha real

    def setUp(self):
        wb = openpyxl.Workbook()
        ws = wb.active
        ws.title = "18.09"
        for c, cab in enumerate(self.CAB, start=1):
            ws.cell(1, c).value = cab
        ws.add_table(Table(displayName="Expedicao", ref="A1:Q5"))
        for c, cab in enumerate(TURNO, start=1):
            ws.cell(10, c).value = cab
        ws.add_table(Table(displayName="Matutino", ref="A10:I11"))
        for c, v in enumerate(["000123", "RODOGRANEL", "JOAO DA SILVA", "ABC1D23", "MG",
                               "XYZ9876", "MG", "LS 4 EIXOS"], start=1):
            ws.cell(11, c).value = v
        caminho = os.path.join(_apoio.pasta_temporaria(self), "EXPED.xlsx")
        wb.save(caminho)
        self.planilha = lan.PlanilhaTeste(caminho)
        self.ws = self.planilha.wb["18.09"]

    def col(self, nome):
        return self.CAB.index(nome) + 1

    def obs(self, linha):
        return self.ws.cell(linha, len(self.CAB)).value

    def sincronizar(self):
        self.planilha.limpar_cache()
        with contextlib.redirect_stdout(io.StringIO()):
            return lan.sincronizar_observacoes(self.planilha, "18.09")

    def test_texto(self):
        self.assertEqual(lan.texto_observacoes(19569.0, "ABC1D23", "xyz-9876", 123456.0),
                         "TICKET N°: 19569 / PLACA CAVALO: ABC-1D23 / PLACA REBOQUE: XYZ-9876 "
                         "/ LACRES N°: 123456")
        self.assertEqual(lan.texto_observacoes(19569, "ABC-1D23", None, "  "),
                         "TICKET N°: 19569 / PLACA CAVALO: ABC-1D23 / PLACA REBOQUE: / LACRES N°:")

    def test_tara_preenche_com_o_lacre_do_turno(self):
        t = {"numero": 19569, "cavalo": "ABC1D23", "reboque": "", "motorista": "JOAO DA SILVA",
             "transportadora": "RODOGRANEL", "entrada": dt.datetime(2026, 9, 18, 7, 4),
             "saida": None, "peso_entrada": 17250.0, "tara": 0.0, "bruto": 0.0, "tipo": "tara"}
        with contextlib.redirect_stdout(io.StringIO()):
            self.assertEqual(lan.lancar_tara(self.planilha, CFG, t)[0], "LANÇADO")
        self.assertEqual(self.sincronizar(), 1)
        self.assertEqual(self.obs(2), "TICKET N°: 19569 / PLACA CAVALO: ABC-1D23 / "
                                      "PLACA REBOQUE: XYZ-9876 / LACRES N°: 000123")
        self.assertEqual(self.sincronizar(), 0)                  # nada mudou: nao regrava

    def test_lacre_digitado_depois_atualiza(self):
        self.ws.cell(2, self.col("TICKET DE PESAGEM")).value = 19569
        self.ws.cell(2, self.col("PLACA DO CAVALO")).value = "ABC-1D23"
        self.assertEqual(self.sincronizar(), 1)
        self.assertTrue(self.obs(2).endswith("LACRES N°:"))
        self.ws.cell(2, self.col("CÓDIGO LACRE")).value = "554433 / 554434"
        self.assertEqual(self.sincronizar(), 1)
        self.assertTrue(self.obs(2).endswith("LACRES N°: 554433 / 554434"))

    def test_nao_mexe_em_texto_escrito_a_mao_nem_em_linha_sem_ticket(self):
        self.ws.cell(2, self.col("TICKET DE PESAGEM")).value = 19569
        self.ws.cell(2, len(self.CAB)).value = "Carga liberada pelo supervisor"
        self.ws.cell(3, self.col("PLACA DO CAVALO")).value = "ABC-1D23"   # sem ticket
        self.assertEqual(self.sincronizar(), 0)
        self.assertEqual(self.obs(2), "Carga liberada pelo supervisor")
        self.assertIsNone(self.obs(3))

    def test_planilha_sem_a_coluna_nao_faz_nada(self):
        wb = openpyxl.Workbook()
        ws = wb.active
        ws.title = "18.09"
        for c, cab in enumerate(EXPED, start=1):
            ws.cell(1, c).value = cab
        ws.add_table(Table(displayName="Expedicao", ref="A1:P2"))
        ws.cell(2, EXPED.index("TICKET DE PESAGEM") + 1).value = 19569
        caminho = os.path.join(_apoio.pasta_temporaria(self), "SEM_X.xlsx")
        wb.save(caminho)
        self.assertEqual(lan.sincronizar_observacoes(lan.PlanilhaTeste(caminho), "18.09"), 0)


if __name__ == "__main__":
    unittest.main()
