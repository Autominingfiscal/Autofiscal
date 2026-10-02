"""Remessas e Faturamento (pacote automining)."""

import datetime as dt
import os
import unittest
from pathlib import Path

import _apoio  # noqa: F401  (ajusta o sys.path)

import openpyxl

from automining import chaves, faturamento, pdf, planilha
from automining.expedicao import nome_eh_data
from automining.planilha import como_numero, inserir_linhas
from automining.refs import deslocar_faixa, deslocar_formula


class TestComoNumero(unittest.TestCase):
    def test_formatos_brasileiros(self):
        self.assertEqual(como_numero("R$ 1.234,56"), 1234.56)
        self.assertEqual(como_numero("14,80"), 14.8)
        self.assertEqual(como_numero("1.234"), 1234.0)
        self.assertEqual(como_numero("14.80"), 14.8)
        self.assertEqual(como_numero(7), 7.0)

    def test_o_que_nao_e_numero(self):
        for v in (None, True, "", "=A1*2", "abc"):
            self.assertIsNone(como_numero(v), v)


class TestRefs(unittest.TestCase):
    def test_estica_faixa_quando_insere_dentro(self):
        self.assertEqual(deslocar_formula("=SUM(C3:C10)", 10, 4, "A", "A"), "=SUM(C3:C14)")

    def test_nao_mexe_depois_do_fim(self):
        self.assertEqual(deslocar_formula("=SUM(C3:C10)", 11, 4, "A", "A"), "=SUM(C3:C10)")

    def test_respeita_aba_e_texto_literal(self):
        f = "='Remessas Porto'!B5&\"B5\"&B5"
        self.assertEqual(deslocar_formula(f, 5, 1, "Remessas Porto", "Outra"),
                         "='Remessas Porto'!B6&\"B5\"&B5")

    def test_nao_confunde_funcao_nem_absoluto(self):
        self.assertEqual(deslocar_formula("=LOG10(A$5)", 2, 1, "A", "A"), "=LOG10(A$6)")

    def test_faixa_solta(self):
        self.assertEqual(deslocar_faixa("A1:C10", 5, 2), "A1:C12")


class TestInserirLinhas(unittest.TestCase):
    def test_formulas_de_outras_abas_acompanham(self):
        wb = openpyxl.Workbook()
        ws = wb.active
        ws.title = "Dados"
        resumo = wb.create_sheet("Resumo")
        resumo["A1"] = "=SUM(Dados!B2:B5)"
        inserir_linhas(wb, ws, 5, 3)
        self.assertEqual(resumo["A1"].value, "=SUM(Dados!B2:B8)")


class TestBackup(unittest.TestCase):
    def test_mesmo_segundo_nao_sobrescreve_e_limpa_as_velhas(self):
        pasta = Path(_apoio.pasta_temporaria(self))
        plan = pasta / "REMESSA.xlsx"
        backup = pasta / "Backup"
        backup.mkdir()
        for dia in range(1, 6):                                  # 5 copias antigas
            (backup / ("2026-09-%02d_080000_REMESSA.xlsx" % dia)).write_text("velha")
        (backup / "2026-09-01_080000_FATURAMENTO.xlsx").write_text("outra planilha")
        (backup / "anotacoes.txt").write_text("nao e backup")

        plan.write_text("v1")
        primeira = planilha.fazer_backup(plan, manter=3)
        plan.write_text("v2")
        segunda = planilha.fazer_backup(plan, manter=3)

        self.assertNotEqual(primeira, segunda)
        self.assertEqual(primeira.read_text(), "v1")             # nao foi sobrescrita
        self.assertEqual(segunda.read_text(), "v2")
        restantes = sorted(p.name for p in backup.iterdir())
        self.assertEqual(len([n for n in restantes if "REMESSA" in n]), 3)
        self.assertIn("2026-09-01_080000_FATURAMENTO.xlsx", restantes)
        self.assertIn("anotacoes.txt", restantes)


class TestChaveDoPdf(unittest.TestCase):
    def test_dv(self):
        chave = _apoio.chave_valida(19007)
        self.assertTrue(pdf.dv_confere(chave))
        errada = chave[:-1] + str((int(chave[-1]) + 1) % 10)
        self.assertFalse(pdf.dv_confere(errada))

    def test_aceita_confere_a_nota(self):
        chave = _apoio.chave_valida(19007)
        self.assertEqual(pdf.aceita(chave, 19007), (True, ""))
        ok, motivo = pdf.aceita(chave, 19008)
        self.assertFalse(ok)
        self.assertIn("19007", motivo)

    def test_le_chave_de_pdf_comprimido_e_sem_compressao(self):
        pasta = _apoio.pasta_temporaria(self)
        chave = _apoio.chave_valida(19007)
        for comprimido in (True, False):
            caminho = os.path.join(pasta, "nota_%s.pdf" % comprimido)
            _apoio.pdf_com_texto(caminho, "CHAVE DE ACESSO", pdf.formatar(chave),
                                 comprimido=comprimido)
            self.assertEqual(pdf.chave_do_pdf(caminho, 19007), (chave, ""))

    def test_formatar(self):
        self.assertEqual(pdf.formatar("1" * 44), " ".join(["1111"] * 11))


class TestIndexarPdfs(unittest.TestCase):
    def test_todos_os_numeros_do_nome(self):
        pasta = Path(_apoio.pasta_temporaria(self))
        (pasta / "2026-09 NF 19007.pdf").write_bytes(b"x")
        (pasta / "sub").mkdir()
        (pasta / "sub" / "NF 19008.PDF").write_bytes(b"x")
        (pasta / "leia.txt").write_text("19009")
        indice, _, arquivos = chaves.indexar_pdfs(pasta, incluir_subpastas=True)
        self.assertIn(19007.0, indice)
        self.assertIn(19008.0, indice)
        self.assertNotIn(19009.0, indice)
        self.assertEqual(arquivos, 2)
        indice, _, _ = chaves.indexar_pdfs(pasta, incluir_subpastas=False)
        self.assertNotIn(19008.0, indice)


class TestExpedicao(unittest.TestCase):
    def test_nome_de_aba(self):
        self.assertTrue(nome_eh_data("17.09"))
        self.assertFalse(nome_eh_data("32.09"))
        self.assertFalse(nome_eh_data("Resumo"))
        self.assertFalse(nome_eh_data("17.09.26"))


class _Cfg:
    uf = "AL"
    cfop_faturamento = 550404.0
    operacao_faturamento = "R.PORTO"
    descricao = "REMESSA"


class TestFaturamento(unittest.TestCase):
    def _mes(self):
        wb = openpyxl.Workbook()
        ws = wb.active
        ws.title = "09-26"
        for r, (dia, nf) in enumerate([(10, 100), (11, 101)], start=3):
            ws.cell(r, 1).value = dt.date(2026, 9, dia)
            ws.cell(r, 1).number_format = "DD/MM/YYYY"
            ws.cell(r, 2).value = nf
            ws.cell(r, 3).value = 10.0
            ws.cell(r, 16).value = "R.PORTO"
        ws.cell(8, 2).value = "SUB-TOTAL"
        ws.cell(8, 3).value = '=SUMIF(A3:A7,">0",C3:C7)'
        return wb, ws

    def test_retroativa_entra_na_ordem_e_formatada(self):
        wb, ws = self._mes()
        avisos = []
        r = faturamento._processar_mes(
            wb, ws, "09-26", {99.0: (dt.date(2026, 9, 5), 5.0, 1.0)}, _Cfg, avisos)
        self.assertEqual(r["inseridas"], 1)
        self.assertEqual([ws.cell(i, 2).value for i in (3, 4, 5)], [99.0, 100, 101])
        self.assertEqual(ws.cell(3, 1).number_format, "DD/MM/YYYY")
        self.assertEqual(avisos, [])

    def test_nota_do_fim_usa_linha_livre(self):
        wb, ws = self._mes()
        r = faturamento._processar_mes(
            wb, ws, "09-26", {102.0: (dt.date(2026, 9, 12), 5.0, 1.0)}, _Cfg, [])
        self.assertEqual(r["inseridas"], 1)
        self.assertEqual(ws.cell(5, 2).value, 102.0)
        self.assertEqual(ws.cell(8, 2).value, "SUB-TOTAL")   # nada foi inserido

    def test_ja_lancada_so_atualiza_valor(self):
        wb, ws = self._mes()
        r = faturamento._processar_mes(
            wb, ws, "09-26", {100.0: (dt.date(2026, 9, 10), 12.5, 1.0)}, _Cfg, [])
        self.assertEqual((r["ja_lancadas"], r["valor_atualizado"]), (1, 1))
        self.assertEqual(ws.cell(3, 3).value, 12.5)


if __name__ == "__main__":
    unittest.main()
