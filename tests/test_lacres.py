"""Lacres do turno: pedir, conferir, gravar na tabela do turno e gerar as etiquetas."""

import contextlib
import datetime as dt
import io
import os
import unittest
from unittest import mock

import _apoio

import openpyxl
from openpyxl.worksheet.table import Table

import etapa3_lancador as lan
import lacres

EXPED = ["TICKET DE PESAGEM", "CÓDIGO LACRE", "MOTORISTA", "PLACA DO CAVALO"]
TURNO = ["PESO A CARREGAR", "CÓDIGO LACRE", "TRANSPORTADORA", "MOTORISTA", "PLACA DO \nCAVALO",
         "PLACA DO REBOQUE", "MODELO"]


class TestLacres(unittest.TestCase):
    def setUp(self):
        wb = openpyxl.Workbook()
        ws = wb.active
        ws.title = "02.10"
        for c, cab in enumerate(EXPED, start=1):
            ws.cell(1, c).value = cab
        ws.add_table(Table(displayName="Expedicao", ref="A1:D3"))
        ws.cell(2, 1).value = 19569                       # ja passou pela balanca
        ws.cell(2, 3).value = "JOSE ADAO"
        ws.cell(2, 4).value = "RNR-9F32"
        for linha_cab, nome in ((10, "Matutino"), (20, "Vespertino")):
            for c, cab in enumerate(TURNO, start=1):
                ws.cell(linha_cab, c).value = cab
            ws.add_table(Table(displayName=nome, ref="A%d:G%d" % (linha_cab, linha_cab + 4)))
        ws.cell(11, 4).value = "MOTORISTA MANHA"
        ws.cell(11, 2).value = "5555 - 6666"
        for i, (nome, cavalo) in enumerate((("PEDRO MELO SILVA", "RMX0I25"),
                                            ("JOSE ADAO", "RNR9F32"),
                                            ("JA TINHA", "AAA1A11"))):
            ws.cell(21 + i, 3).value = "RODOGRANEL"
            ws.cell(21 + i, 4).value = nome
            ws.cell(21 + i, 5).value = cavalo
            ws.cell(21 + i, 6).value = "RNG6B29"
            ws.cell(21 + i, 7).value = "vanderleia"
        ws.cell(23, 2).value = "1111 - 2222"
        self.caminho = os.path.join(_apoio.pasta_temporaria(self), "EXPED.xlsx")
        wb.save(self.caminho)
        self.planilha = lan.PlanilhaTeste(self.caminho)
        self.ws = self.planilha.wb["02.10"]

    def rodar(self, respostas):
        fila = list(respostas)
        perguntas = []

        def perguntar(txt):
            perguntas.append(txt)
            if not fila:
                raise EOFError
            return fila.pop(0)

        tab, motoristas = lacres.motoristas_do_turno(self.planilha, "02.10", 2)
        saida = io.StringIO()
        with contextlib.redirect_stdout(saida):
            n = lacres.registrar(self.planilha, "02.10", tab, motoristas, perguntar)
        return n, motoristas, perguntas, saida.getvalue()

    def test_numeros(self):
        self.assertEqual(lacres.numeros("9999 - 1265"), ["9999", "1265"])
        self.assertEqual(lacres.numeros("9999/1265"), ["9999", "1265"])
        self.assertEqual(lacres.numeros(1131121.0), ["1131121"])
        self.assertEqual(lacres.numeros(None), [])

    def test_so_pergunta_quem_nao_tem_e_confere(self):
        n, motoristas, perguntas, saida = self.rodar(
            ["9999 1265",          # PEDRO
             "5555 7777",          # JOSE ADAO: 5555 ja e do matutino
             "1835",               # so um numero
             "1835 1835",          # o mesmo duas vezes
             "1111 3333",          # 1111 ja e do JA TINHA (vespertino)
             "1835-1743"])
        self.assertEqual(n, 2)
        self.assertEqual(len([p for p in perguntas if "JA TINHA" in p]), 0)
        self.assertIn("ja esta em aba 02.10, MOTORISTA MANHA", saida)
        self.assertIn("Li 1: 1835", saida)
        self.assertIn("mesmo numero", saida)
        self.assertIn("JA TINHA", saida)                       # o 1111 em uso
        self.assertEqual(self.ws.cell(21, 2).value, "9999 - 1265")
        self.assertEqual(self.ws.cell(22, 2).value, "1835 - 1743")
        self.assertEqual(self.ws.cell(23, 2).value, "1111 - 2222")   # nao mexeu
        # JOSE ADAO ja tinha passado pela balanca: a linha da expedicao ganha o lacre
        self.assertEqual(self.ws.cell(2, 2).value, "1835 - 1743")

    def test_enter_vazio_pula(self):
        n, motoristas, _, saida = self.rodar(["", "1835 1743"])
        self.assertEqual(n, 1)
        self.assertIsNone(self.ws.cell(21, 2).value)
        self.assertIn("pulado: PEDRO MELO SILVA", saida)

    def test_etiquetas_15_por_folha(self):
        _, motoristas = lacres.motoristas_do_turno(self.planilha, "02.10", 2)
        m = dict(motoristas[0], lacres=["9999", "1265"])
        html = lacres.etiquetas_html([m] * 16 + [dict(m, lacres=[])], "vespertino",
                                     dt.date(2026, 10, 2))
        self.assertEqual(html.count('class="folha"'), 2)
        self.assertEqual(html.count('class="etiqueta"'), 16)       # sem lacre nao sai etiqueta
        for trecho in ("PEDRO MELO SILVA", "RMX-0I25", "RNG-6B29", "VANDERLEIA", "RODOGRANEL",
                       "9999 – 1265", "Vespertino · 02/10/2026"):
            self.assertIn(trecho, html)

    def test_main_em_teste_nao_salva_e_gera_etiqueta(self):
        destino = _apoio.pasta_temporaria(self)
        respostas = iter(["9999 1265", "1835 1743"])
        with mock.patch.object(lacres, "PASTA_ETIQUETAS", destino), \
                contextlib.redirect_stdout(io.StringIO()):
            codigo = lacres.main(["--turno", "vespertino", "--aba", "02.10", "--teste",
                                  self.caminho, "--nao-abrir"], lambda txt: next(respostas))
        self.assertEqual(codigo, 0)
        self.assertEqual(os.listdir(destino), ["lacres_2026-10-02_vespertino.html"])
        self.assertIsNone(openpyxl.load_workbook(self.caminho)["02.10"].cell(21, 2).value)

    def test_data_da_aba(self):
        self.assertEqual(lacres.data_da_aba("02.10", dt.date(2026, 10, 5)), dt.date(2026, 10, 2))
        self.assertEqual(lacres.data_da_aba("28.12", dt.date(2027, 1, 3)), dt.date(2026, 12, 28))


if __name__ == "__main__":
    unittest.main()
