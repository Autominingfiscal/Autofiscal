# -*- coding: utf-8 -*-
"""
CONFERÊNCIA: RETENÇÕES ISSQN  x  2-CONFERENCIA RETENÇÕES 2026

COMO USAR
  Dê dois cliques em conferir_issqn.py (ou renomeie para .pyw para não
  aparecer a janela preta).
  1. Escolha a planilha "Retenções ISSQN" do mês.
  2. Escolha a planilha "2-CONFERENCIA RETENÇÕES".
     A aba é escolhida pelo mês do nome do arquivo ISSQN
     (ex.: ..._092026.xlsx -> aba 09_2026). Se não achar, pergunta.
  3. Escolha a pasta onde as notas ficam salvas (a última fica lembrada).
     Cancelar = rodar sem verificar a pasta.
  4. O resultado é salvo como uma planilha NOVA na mesma pasta deste script
     e abre sozinho. As planilhas originais não são alteradas.
  Para não precisar escolher tudo toda vez, preencha os caminhos no bloco
  CONFIGURAÇÃO logo abaixo.

  Também roda pelo Prompt de Comando, sem janelas:
     python conferir_issqn.py ISSQN.xlsx CONFERENCIA.xlsx "C:\\pasta\\das\\notas"
     (opcional: --aba 09_2026   --sem-abrir)

O QUE COMPARA
  Só entram notas COM ISSQN:
  Planilha ISSQN : Razão Social, Nr. da NF, ISSQN  -> só linhas com ISSQN preenchido
  Conferência    : FORNECEDOR, TÍTULO, coluna Q (ISS 2090)
                   -> só as linhas com valor preenchido na coluna Q entram
                      na comparação; as demais são ignoradas.
  - Casa a nota pelo número (NF x Título), ignorando zeros à esquerda,
    letras e traços. O nome do fornecedor desempata e é conferido
    (ignora LTDA, ME, EPP, EIRELI, S/A, acentos e pontuação).
  - Valor: compara o ISSQN da Retenções com a coluna Q; diferença de até
    1 centavo conta como arredondamento.
  - Pasta: procura arquivo cujo NOME contenha o número da NF, e lista os
    arquivos da pasta que não batem com nenhuma nota da Retenções nem da
    conferência (os de notas sem ISS na conferência são ignorados).

IMPORTANTE
  O Python lê os valores que ficaram SALVOS nas planilhas. Se alterou algo
  no Excel, salve antes de rodar.
"""

import argparse
import datetime
import json
import os
import re
import sys
import traceback
import unicodedata

from openpyxl import Workbook, load_workbook
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.utils import column_index_from_string

# ---------------- CONFIGURAÇÃO (pode alterar) ----------------
# Cole os caminhos entre as aspas, depois do r. Deixe r"" para escolher
# numa janela na hora. Dica: no Windows Explorer, clique com o botão direito
# no arquivo/pasta segurando Shift > "Copiar como caminho" (tire as aspas).

# Pasta onde ficam os PDFs das notas.
# Ex.: PASTA_NOTAS = r"C:\Fiscal\Notas 2026"   ou   r"\\servidor\fiscal\NFs 2026"
PASTA_NOTAS = r""

# Pasta onde ficam as planilhas "Retenções ISSQN" (a janela já abre nela;
# como o arquivo muda todo mês, ele continua perguntando qual é).
# Ex.: PASTA_PLANILHAS_ISSQN = r"C:\Fiscal\ISSQN\2026"
PASTA_PLANILHAS_ISSQN = r""

# Arquivo da conferência (é o mesmo o ano todo). Preenchido = não pergunta.
# Ex.: PLANILHA_CONFERENCIA = r"C:\Fiscal\2-CONFERENCIA RETENÇÕES 2026 atual.xlsx"
PLANILHA_CONFERENCIA = r""
INCLUIR_SUBPASTAS = True        # procurar também nas subpastas
COLUNA_ISS_CONF = "Q"           # coluna do ISS na conferência
TOLERANCIA = 0.01               # R$ 0,01
EXTENSOES_NOTAS = None          # ex.: {".pdf", ".xml"}  |  None = qualquer arquivo
# --------------------------------------------------------------

PASTA_DO_SCRIPT = os.path.dirname(os.path.abspath(__file__))
ARQ_CONFIG = os.path.join(PASTA_DO_SCRIPT, "conferir_issqn_config.json")
MAX_COL = 60

VERDE = PatternFill("solid", start_color="C6EFCE")
AMARELO = PatternFill("solid", start_color="FFEB9C")
VERMELHO = PatternFill("solid", start_color="FFC7CE")
AZUL = PatternFill("solid", start_color="1F4E79")
FORMATO_RS = '"R$" #,##0.00'

_tk_root = None


# ============================================================
#  Texto, números e nomes
# ============================================================

def normalizar(s):
    """Maiúsculas, sem acento, só letras/números separados por um espaço."""
    if s is None:
        return ""
    s = str(s).replace("&", " E ")          # "&" vale o mesmo que "e"
    s = unicodedata.normalize("NFKD", s)
    s = "".join(c for c in s if not unicodedata.combining(c))
    s = re.sub(r"[^A-Z0-9]+", " ", s.upper())
    return s.strip()


IGNORAR_NOME = {"LTDA", "LTD", "ME", "EPP", "EIRELI", "SA", "S", "A", "CIA", "SLU",
                "MEI", "DE", "DA", "DO", "DOS", "DAS", "E", "EM"}


def nome_limpo(s):
    return " ".join(w for w in normalizar(s).split() if w not in IGNORAR_NOME)


def comparar_nomes(a, b):
    """2 = igual ou um contém o outro | 1 = parecido | 0 = diferente."""
    if not a or not b:
        return 0
    if a == b or f" {b} " in f" {a} " or f" {a} " in f" {b} ":
        return 2
    pa, pb = a.split(), b.split()
    if pa[0] == pb[0]:
        return 1
    curto, longo = (pa, pb) if len(a) <= len(b) else (pb, pa)
    comuns = sum(1 for w in curto if w in longo)
    return 1 if curto and comuns / len(curto) >= 0.6 else 0


def numeros_do_texto(s):
    """Todas as sequências de dígitos, sem zeros à esquerda."""
    return [n.lstrip("0") or "0" for n in re.findall(r"\d+", s or "")]


def numero_principal(s):
    """Maior sequência de dígitos ("000123-1" -> "123")."""
    nums = numeros_do_texto(s)
    return max(nums, key=len) if nums else ""


def numeros_combinam(a, b):
    """2 = mesmo número | 1 = um termina com o outro, separado por zero
    (ex.: 202600000000123 x 123) | 0 = diferente."""
    if not a or not b:
        return 0
    if a == b:
        return 2
    lg, ct = (a, b) if len(a) > len(b) else (b, a)
    if len(ct) >= 3 and lg.endswith(ct) and lg[-len(ct) - 1] == "0":
        return 1
    return 0


def texto_cel(v):
    if v is None:
        return ""
    if isinstance(v, str) and v.startswith("#"):   # #REF!, #N/D...
        return ""
    return str(v).strip()


def texto_nf(v):
    if v is None:
        return ""
    if isinstance(v, bool):
        return ""
    if isinstance(v, float) and v.is_integer():
        return str(int(v))
    if isinstance(v, (int, float)):
        return str(v)
    return texto_cel(v)


def valor_num(v):
    """Retorna (valor, tem_valor)."""
    if v is None or isinstance(v, bool):
        return 0.0, False
    if isinstance(v, (int, float)):
        return float(v), True
    s = str(v).replace("R$", "").replace("\xa0", "").strip()
    if s in ("", "-") or s.startswith("#"):
        return 0.0, False
    if "," in s:                      # formato brasileiro 1.234,56
        s = s.replace(".", "").replace(",", ".")
    try:
        return float(s), True
    except ValueError:
        return 0.0, False


# ============================================================
#  Leitura das planilhas
# ============================================================

def ler_linhas(ws):
    """Lê a aba inteira (até MAX_COL colunas) como lista de listas, índice 0 = linha 1."""
    linhas = []
    for row in ws.iter_rows(min_row=1, max_col=MAX_COL, values_only=True):
        row = list(row) + [None] * (MAX_COL - len(row))
        linhas.append(row)
    return linhas


def achar_cabecalho(linhas, alvo):
    """Procura nas 30 primeiras linhas uma célula igual a `alvo` (normalizado)."""
    for r, row in enumerate(linhas[:30]):
        for c, v in enumerate(row):
            if normalizar(texto_cel(v)) == alvo:
                return r, c
    return None, None


def achar_coluna_palavra(row, palavra):
    for c, v in enumerate(row):
        if f" {palavra} " in f" {normalizar(texto_cel(v))} ":
            return c
    return None


def aba_pelo_nome_arquivo(nome):
    """'..._092026.xlsx' -> '09_2026'."""
    for m in re.finditer(r"(?=(\d{6}))", nome):
        t = m.group(1)
        if 1 <= int(t[:2]) <= 12 and t[2:4] == "20":
            return f"{t[:2]}_{t[2:]}"
    return ""


def achar_aba(wb, nome):
    if not nome or not nome.strip():
        return None
    for ws in wb.worksheets:
        if ws.title.strip().lower() == nome.strip().lower():
            return ws
    return None


def abrir(caminho):
    ext = os.path.splitext(caminho)[1].lower()
    if ext == ".xls":
        raise RuntimeError(f"'{os.path.basename(caminho)}' está no formato antigo .xls. "
                           "Abra no Excel e salve como .xlsx.")
    return load_workbook(caminho, read_only=True, data_only=True)


def ler_issqn(caminho):
    wb = abrir(caminho)
    try:
        for ws in wb.worksheets:
            linhas = ler_linhas(ws)
            lin_cab, c_razao = achar_cabecalho(linhas, "RAZAO SOCIAL")
            if lin_cab is not None:
                break
        else:
            raise RuntimeError("Não encontrei o cabeçalho 'Razão Social' na planilha ISSQN.")
        cab = linhas[lin_cab]
        c_nf = achar_coluna_palavra(cab, "NF")
        c_iss = achar_coluna_palavra(cab, "ISSQN")
        if c_nf is None or c_iss is None:
            raise RuntimeError("Não encontrei as colunas 'Nr. da NF' e/ou 'ISSQN' na planilha ISSQN.")

        notas, nums_sem_iss = [], set()
        for idx in range(lin_cab + 1, len(linhas)):
            row = linhas[idx]
            razao = texto_cel(row[c_razao])
            if normalizar(razao).startswith("BASE DE CALCULO"):
                break
            nf_txt = texto_nf(row[c_nf])
            if not razao and not nf_txt:
                continue
            valor, _ = valor_num(row[c_iss])
            if abs(valor) < 0.005:               # só notas com ISSQN preenchido
                if numero_principal(nf_txt):
                    nums_sem_iss.add(numero_principal(nf_txt))
                continue
            notas.append({
                "linha": idx + 1, "razao": razao, "razao_n": nome_limpo(razao),
                "nf_txt": nf_txt, "nf": numero_principal(nf_txt), "valor": valor,
            })
        return notas, ws.title, nums_sem_iss
    finally:
        wb.close()


def ler_conferencia(caminho, nome_arq_issqn, aba_escolhida=None, perguntar_aba=None):
    wb = abrir(caminho)
    try:
        nome_aba = aba_escolhida or aba_pelo_nome_arquivo(nome_arq_issqn)
        ws = achar_aba(wb, nome_aba)
        if ws is None:
            abas = [w.title.strip() for w in wb.worksheets]
            if perguntar_aba is None:
                raise RuntimeError(f"Aba '{nome_aba}' não encontrada. Abas: {', '.join(abas)}. "
                                   "Use --aba para escolher.")
            nome_aba = perguntar_aba(abas, nome_aba)
            if not nome_aba:
                return None, None, None
            ws = achar_aba(wb, nome_aba)
            if ws is None:
                raise RuntimeError(f"Aba '{nome_aba}' não encontrada.")

        linhas = ler_linhas(ws)
        lin_cab, c_forn = achar_cabecalho(linhas, "FORNECEDOR")
        if lin_cab is None:
            raise RuntimeError(f"Não encontrei o cabeçalho 'FORNECEDOR' na aba {ws.title}.")
        c_tit = achar_coluna_palavra(linhas[lin_cab], "TITULO")
        if c_tit is None:
            raise RuntimeError(f"Não encontrei a coluna 'TÍTULO' na aba {ws.title}.")
        c_q = column_index_from_string(COLUNA_ISS_CONF) - 1

        conf, nums_sem_q = [], set()
        for idx in range(lin_cab + 1, len(linhas)):
            row = linhas[idx]
            if any(normalizar(texto_cel(v)) == "TOTAL" for v in row[:12]):
                break
            forn = texto_cel(row[c_forn])
            tit = texto_nf(row[c_tit])
            valor, tem = valor_num(row[c_q])
            if not tem or abs(valor) < 0.005:      # só linhas com valor preenchido na coluna Q
                if numero_principal(tit):
                    nums_sem_q.add(numero_principal(tit))
                continue
            conf.append({
                "linha": idx + 1, "forn": forn, "forn_n": nome_limpo(forn),
                "tit_txt": tit, "tit": numero_principal(tit),
                "valor": valor, "tem": tem, "usado": False,
            })
        return conf, ws.title.strip(), nums_sem_q
    finally:
        wb.close()


# ============================================================
#  Pasta das notas
# ============================================================

def listar_arquivos(pasta):
    arquivos, indice = [], {}
    if INCLUIR_SUBPASTAS:
        caminhos = (os.path.join(raiz, f) for raiz, _, fs in os.walk(pasta) for f in fs)
    else:
        caminhos = (os.path.join(pasta, f) for f in os.listdir(pasta)
                    if os.path.isfile(os.path.join(pasta, f)))
    for caminho in caminhos:
        nome = os.path.basename(caminho)
        if nome.startswith("~$") or nome.startswith("Resultado conferência ISSQN"):
            continue
        if EXTENSOES_NOTAS and os.path.splitext(nome)[1].lower() not in EXTENSOES_NOTAS:
            continue
        nums = set(numeros_do_texto(os.path.splitext(nome)[0]))
        item = {"caminho": caminho, "nome": nome, "nome_n": f" {normalizar(nome)} ", "nums": nums}
        arquivos.append(item)
        for n in nums:
            indice.setdefault(n, []).append(item)
    return arquivos, indice


PALAVRAS_GENERICAS_ARQ = {"NOTA", "NOTAS", "FISCAL", "NFSE", "NFSEN", "DANFE", "DANFSE", "SERVICO",
                          "SERVICOS", "PDF", "XML", "JPG", "JPEG", "PNG", "COPIA", "CÓPIA", "ISSQN",
                          "RETENCAO", "RECIBO", "BOLETO", "FATURA", "CNPJ", "EMISSAO"}


def buscar_arquivo(nf, nome_n, arquivos, indice, tem_pasta):
    """Retorna (status, arquivo_escolhido, quantidade_de_candidatos)."""
    if not tem_pasta:
        return "PASTA NÃO INFORMADA", None, 0
    if not nf:
        return "SEM Nº DE NF", None, 0
    candidatos = indice.get(nf, [])
    exato = bool(candidatos)
    if not exato:
        candidatos = [a for a in arquivos if any(numeros_combinam(nf, t) == 1 for t in a["nums"])]
    if not candidatos:
        return "NÃO", None, 0
    palavras = [w for w in nome_n.split() if len(w) >= 3]
    def bate_nome(a):
        return any(f" {w} " in a["nome_n"] for w in palavras)
    escolhido = next((a for a in candidatos if bate_nome(a)), candidatos[0])
    if not exato:
        return "SIM (nº parcial - conferir)", escolhido, len(candidatos)
    # arquivo com o número certo, mas com nome de outra empresa escrito nele
    palavras_arq = [w for w in escolhido["nome_n"].split()
                    if len(w) >= 4 and not w.isdigit() and w not in PALAVRAS_GENERICAS_ARQ]
    if palavras and palavras_arq and not bate_nome(escolhido):
        return "SIM (outro fornecedor? conferir)", escolhido, len(candidatos)
    return "SIM", escolhido, len(candidatos)


# ============================================================
#  Comparação
# ============================================================

def casar(nota, conf):
    """Escolhe a melhor linha da conferência para a nota.
    Retorna (linha, tipo, outro): `outro` é uma linha com o MESMO número mas de
    outro fornecedor (não conta como a nota, só serve de aviso)."""
    melhor, melhor_pts, tipo, outro = None, -1, 0, None
    for c in conf:
        sn = numeros_combinam(nota["nf"], c["tit"])
        snome = comparar_nomes(nota["razao_n"], c["forn_n"])
        ok_val = abs(c["valor"] - nota["valor"]) <= TOLERANCIA
        if sn == 2 and snome == 0 and c["forn_n"]:
            outro = outro or c          # mesmo nº, outro fornecedor: não é a mesma nota
            continue
        if sn == 2:
            pts = 100 + snome * 10
        elif sn == 1 and snome >= 1 and not c["usado"]:
            pts = 60 + snome * 10
        elif snome >= 1 and ok_val and nota["valor"] > 0 and not c["usado"]:
            pts = 20 + snome * 10
        else:
            continue
        pts += 5 if ok_val else 0
        pts += 3 if not c["usado"] else 0
        if pts > melhor_pts:
            melhor, melhor_pts, tipo = c, pts, sn
    return melhor, tipo, outro


def st_pasta(st_arq, tem_pasta):
    if not tem_pasta:
        return "-"
    if st_arq == "SIM":
        return "SIM"
    if st_arq.startswith("SIM"):
        return "CONFERIR"
    return "NÃO"


def o_que_falta(item, tem_pasta):
    """Texto e categoria a partir de SIM/NÃO/CONFERIR dos três lugares."""
    lugares = (("Retenções", item["na_ret"]), ("conferência", item["na_conf"]), ("pasta", item["na_pasta"]))
    faltam = [nome for nome, st in lugares if st == "NÃO"]
    duvida = [nome for nome, st in lugares if st == "CONFERIR"]
    valor_ruim = item["valor_st"] == "NÃO"
    partes = []
    if faltam:
        partes.append("Falta na " + " e na ".join(faltam))
    if valor_ruim:
        dif = item["v_ret"] - item["v_conf"]
        partes.append(f"Valor não bate: Retenções {moeda(item['v_ret'])} x conferência "
                      f"{moeda(item['v_conf'])} (diferença {moeda(dif)})")
    if duvida:
        partes.append("Conferir " + " e ".join(duvida))
    if partes:
        return ". ".join(partes), ("falta" if faltam or valor_ruim else "conferir")
    return "Nada - nota completa e valor batendo" + ("" if tem_pasta else " (pasta não verificada)"), "completa"


def numero_do_arquivo(nums):
    """Número mais provável de NF no nome do arquivo (evita CNPJ/datas longas)."""
    curtos = [n for n in nums if len(n) <= 10]
    return max(curtos or list(nums), key=len) if nums else ""


def conferir(notas, conf, arquivos, indice, tem_pasta, nome_aba, nums_sem_iss=frozenset()):
    """Uma linha por nota COM ISSQN: as da planilha Retenções (com valor no ISSQN)
    e as da conferência (com valor na coluna Q) que não estão na Retenções."""
    itens = []

    # 1) notas da planilha Retenções
    for nota in notas:
        c, tipo, outro = casar(nota, conf)
        obs = []
        na_conf = "SIM"
        if c is None:
            na_conf = "NÃO"
            if outro:
                obs.append(f"Na conferência existe o título {outro['tit_txt']} com o mesmo número, "
                           f"mas do fornecedor {outro['forn']} (linha {outro['linha']}).")
        else:
            if c["usado"]:
                na_conf = "NÃO"
                obs.append(f"A linha {c['linha']} da conferência já foi usada por outra nota: "
                           "esta pode estar faltando na conferência ou duplicada na Retenções.")
            elif tipo == 0:
                na_conf = "CONFERIR"
                obs.append(f"O nº da NF não aparece na conferência. Pode ser o título {c['tit_txt']} "
                           "(mesmo fornecedor e valor) lançado com número errado.")
            elif tipo == 1:
                obs.append(f"Número casado de forma parcial com o título {c['tit_txt']} (zeros/ano).")
            if c["forn_n"] and comparar_nomes(nota["razao_n"], c["forn_n"]) == 0:
                obs.append(f"Na conferência o fornecedor está como: {c['forn']}.")
            c["usado"] = True

        st_arq, arq, qtd = buscar_arquivo(nota["nf"], nota["razao_n"], arquivos, indice, tem_pasta)
        if st_arq.startswith("SIM ("):
            obs.append("Arquivo: " + st_arq[st_arq.find("(") + 1:-1] + ".")
        if qtd > 1:
            obs.append(f"{qtd} arquivos com esse número na pasta.")

        conf_ok = c if na_conf != "NÃO" else None
        if conf_ok is None:
            valor_st = "-"
        elif abs(nota["valor"] - conf_ok["valor"]) <= TOLERANCIA:
            valor_st = "SIM"
        else:
            valor_st = "NÃO"
        itens.append({
            "nome": nota["razao"], "nf_txt": nota["nf_txt"],
            "v_ret": nota["valor"], "v_conf": conf_ok["valor"] if conf_ok else None,
            "lin_ret": nota["linha"], "lin_conf": conf_ok["linha"] if conf_ok else None,
            "na_ret": "SIM", "na_conf": na_conf, "na_pasta": st_pasta(st_arq, tem_pasta),
            "valor_st": valor_st, "arq": arq, "obs": " ".join(obs), "origem": "ret",
        })

    # 2) notas com ISS na conferência (col. Q) que não estão na Retenções
    for c in conf:
        if c["usado"]:
            continue
        obs = []
        mesmo_num = next((n for n in notas if n["nf"] and numeros_combinam(n["nf"], c["tit"]) == 2), None)
        if mesmo_num:
            obs.append(f"Na Retenções existe a NF {mesmo_num['nf_txt']} com o mesmo número, "
                       f"mas de {mesmo_num['razao']} (linha {mesmo_num['linha']}).")
        st_arq, arq, qtd = buscar_arquivo(c["tit"], c["forn_n"], arquivos, indice, tem_pasta)
        if st_arq.startswith("SIM ("):
            obs.append("Arquivo: " + st_arq[st_arq.find("(") + 1:-1] + ".")
        if qtd > 1:
            obs.append(f"{qtd} arquivos com esse número na pasta.")
        itens.append({
            "nome": c["forn"], "nf_txt": c["tit_txt"],
            "v_ret": None, "v_conf": c["valor"],
            "lin_ret": None, "lin_conf": c["linha"],
            "na_ret": "NÃO", "na_conf": "SIM", "na_pasta": st_pasta(st_arq, tem_pasta),
            "valor_st": "-", "arq": arq, "obs": " ".join(obs), "origem": "conf",
        })

    # 3) arquivos na pasta que não correspondem a nenhuma nota da Retenções nem da conferência
    usados = {id(i["arq"]) for i in itens if i["arq"]}
    nums_iss = {n["nf"] for n in notas if n["nf"]} | {c["tit"] for c in conf if c["tit"]}
    ignorados = 0
    for a in arquivos:
        if id(a) in usados or not a["nums"] or a["nums"] & nums_iss:
            continue
        if a["nums"] & nums_sem_iss:
            ignorados += 1          # nota sem ISS (na conferência sem valor na Q, ou ISSQN zerado)
            continue
        itens.append({
            "nome": os.path.splitext(a["nome"])[0], "nf_txt": numero_do_arquivo(a["nums"]),
            "v_ret": None, "v_conf": None, "lin_ret": None, "lin_conf": None,
            "na_ret": "NÃO", "na_conf": "NÃO", "na_pasta": "SIM",
            "valor_st": "-", "arq": a, "origem": "pasta",
            "obs": "Arquivo na pasta sem nota correspondente na Retenções nem na conferência (col. "
                   f"{COLUNA_ISS_CONF}). Fornecedor e número tirados do nome do arquivo.",
        })

    for it in itens:
        it["falta"], it["cat"] = o_que_falta(it, tem_pasta)
    return itens, ignorados


# ============================================================
#  Planilha de resultado
# ============================================================

def cor_status(cel):
    s = str(cel.value or "")
    if s in ("OK", "SIM"):
        cel.fill = VERDE
    elif s == "CONFERIR":
        cel.fill = AMARELO
    elif s == "NÃO":
        cel.fill = VERMELHO


def formatar_tabela(ws, larguras, colunas_rs):
    for cel in ws[1]:
        cel.font = Font(name="Arial", bold=True, color="FFFFFF")
        cel.fill = AZUL
        cel.alignment = Alignment(wrap_text=True, vertical="center", horizontal="center")
    for row in ws.iter_rows(min_row=2):
        for cel in row:
            cel.font = Font(name="Arial", size=10, color=cel.font.color, underline=cel.font.underline,
                            bold=cel.font.bold)
    for letra, larg in larguras.items():
        ws.column_dimensions[letra].width = larg
    for letra in colunas_rs:
        for cel in ws[letra][1:]:
            cel.number_format = FORMATO_RS
    ws.freeze_panes = "C2"
    ws.auto_filter.ref = ws.dimensions


def link(cel, arq):
    cel.value = arq["nome"]
    cel.hyperlink = arq["caminho"]
    cel.font = Font(name="Arial", size=10, color="0563C1", underline="single")


def gerar_resultado(itens, conf, info, ignorados=0):
    wb = Workbook()
    wsres = wb.active
    wsres.title = "Resumo"
    wsc = wb.create_sheet("Comparação")
    tem_pasta = info["pasta"] is not None

    # --------------------- Comparação ---------------------
    wsc.append(["Fornecedor / arquivo", "Nº NF / Título", "ISSQN Retenções", f"ISS Conferência (col. {COLUNA_ISS_CONF})",
                "Diferença", "Na Retenções", "Na Conferência", "Na Pasta", "Valor bate?",
                "O que falta / o que não bate", "Arquivo", "Linha Retenções", "Linha Conferência", "Observação"])
    for r, it in enumerate(itens, start=2):
        wsc.cell(r, 1, it["nome"])
        wsc.cell(r, 2, it["nf_txt"]).number_format = "@"
        wsc.cell(r, 3, it["v_ret"])
        wsc.cell(r, 4, it["v_conf"])
        if it["v_ret"] is not None and it["v_conf"] is not None:
            wsc.cell(r, 5, f"=C{r}-D{r}")
        wsc.cell(r, 6, it["na_ret"])
        wsc.cell(r, 7, it["na_conf"])
        wsc.cell(r, 8, it["na_pasta"])
        wsc.cell(r, 9, it["valor_st"])
        wsc.cell(r, 10, it["falta"])
        if it["arq"]:
            link(wsc.cell(r, 11), it["arq"])
        wsc.cell(r, 12, it["lin_ret"])
        wsc.cell(r, 13, it["lin_conf"])
        wsc.cell(r, 14, it["obs"] or None)
        wsc.cell(r, 10).font = Font(name="Arial", size=10, bold=it["cat"] != "completa")
    if not itens:
        wsc.cell(2, 1, "Nenhuma nota com ISSQN encontrada nas planilhas.")
    formatar_tabela(wsc, {"A": 38, "B": 16, "C": 15, "D": 17, "E": 13, "F": 12, "G": 13, "H": 11,
                          "I": 12, "J": 55, "K": 32, "L": 11, "M": 12, "N": 60}, "CDE")
    for r, it in enumerate(itens, start=2):
        for col in (6, 7, 8, 9):
            cor_status(wsc.cell(r, col))
            wsc.cell(r, col).alignment = Alignment(horizontal="center")
        wsc.cell(r, 10).fill = {"completa": VERDE, "conferir": AMARELO}.get(it["cat"], VERMELHO)
        wsc.cell(r, 14).alignment = Alignment(wrap_text=True, vertical="top")

    # ----------------------- Resumo -----------------------
    ult = max(len(itens) + 1, 2)
    notas_iss = [i for i in itens if i["origem"] != "pasta"]
    cont = {
        "total": len(notas_iss),
        "completa": sum(1 for i in itens if i["cat"] == "completa"),
        "f_ret": sum(1 for i in notas_iss if i["na_ret"] == "NÃO"),
        "f_conf": sum(1 for i in notas_iss if i["na_conf"] == "NÃO"),
        "so_pasta": sum(1 for i in itens if i["origem"] == "pasta"),
        "ignorados": ignorados,
        "f_pasta": sum(1 for i in itens if i["na_pasta"] == "NÃO"),
        "conferir": sum(1 for i in itens if "CONFERIR" in (i["na_conf"], i["na_pasta"])),
        "valor": sum(1 for i in itens if i["valor_st"] == "NÃO"),
        "valor_ok": sum(1 for i in itens if i["valor_st"] == "SIM"),
    }
    total_q = sum(c["valor"] for c in conf)
    linhas = [
        ("titulo", "Conferência de notas com ISSQN", None),
        (None, None, None),
        (None, "Gerado em", datetime.datetime.now().strftime("%d/%m/%Y %H:%M")),
        (None, "Planilha Retenções ISSQN", f'{info["arq_issqn"]}  (aba {info["aba_issqn"]})'),
        (None, "Planilha de conferência", f'{info["arq_conf"]}  (aba {info["aba_conf"]}, col. {COLUNA_ISS_CONF})'),
        (None, "Pasta das notas", f'{info["pasta"]}  -  {info["n_arq"]} arquivos' if tem_pasta else "(não verificada)"),
        (None, None, None),
        ("sec", "NOTAS COM ISSQN", None),
        ("total", "Total de notas com ISSQN (Retenções + conferência)", cont["total"]),
        ("completa", "Completas (nos 3 lugares e com valor batendo)", cont["completa"]),
        ("f_ret", "Faltam na planilha Retenções", cont["f_ret"]),
        ("f_conf", "Faltam na planilha de conferência", cont["f_conf"]),
        ("f_pasta", "Faltam na pasta", cont["f_pasta"]),
        ("conferir", "Para conferir (número/arquivo duvidoso)", cont["conferir"]),
        ("valor_ok", f"Valor bate (Retenções = col. {COLUNA_ISS_CONF})", cont["valor_ok"]),
        ("valor", f"Valor NÃO bate (Retenções x col. {COLUNA_ISS_CONF})", cont["valor"]),
        (None, None, None),
        ("sec", "ARQUIVOS DA PASTA", None),
        ("so_pasta", "Arquivos na pasta que NÃO estão na Retenções nem na conferência",
         cont["so_pasta"] if tem_pasta else "(pasta não verificada)"),
        ("ignorados", f"Arquivos de notas sem ISS (conferência sem valor na col. {COLUNA_ISS_CONF}) - ignorados",
         cont["ignorados"] if tem_pasta else "-"),
        (None, None, None),
        ("sec", "VALORES", None),
        ("t_ret", "Total ISSQN - planilha Retenções", f"=SUM('Comparação'!C2:C{ult})"),
        ("t_q", f"Total ISS - conferência (col. {COLUNA_ISS_CONF})", total_q),
        ("t_dif", "Diferença", None),
    ]
    pos = {}
    for r, (chave, a, b) in enumerate(linhas, start=1):
        if chave and chave != "sec":
            pos[chave] = r
        ca, cb = wsres.cell(r, 1, a), wsres.cell(r, 2, b)
        ca.font = Font(name="Arial", bold=r >= 8, size=10)
        cb.font = Font(name="Arial", size=10)
        if chave == "sec":
            ca.font = Font(name="Arial", bold=True, size=10, color="FFFFFF")
            ca.fill = cb.fill = AZUL
        elif r >= 8:
            cb.alignment = Alignment(horizontal="right")
    wsres["A1"].font = Font(name="Arial", bold=True, size=14)
    wsres.cell(pos["t_dif"], 2, f"=B{pos['t_ret']}-B{pos['t_q']}")
    for k in ("t_ret", "t_q", "t_dif"):
        wsres.cell(pos[k], 2).number_format = FORMATO_RS
        wsres.cell(pos[k], 2).alignment = Alignment(horizontal="right")
    if cont["total"]:
        for col in (1, 2):
            wsres.cell(pos["completa"], col).fill = VERDE
            wsres.cell(pos["valor_ok"], col).fill = VERDE
    for k, cor in (("f_ret", VERMELHO), ("f_conf", VERMELHO), ("f_pasta", VERMELHO),
                   ("conferir", AMARELO), ("valor", VERMELHO), ("so_pasta", VERMELHO)):
        if cont[k]:
            for col in (1, 2):
                wsres.cell(pos[k], col).fill = cor
    wsres.column_dimensions["A"].width = 52
    wsres.column_dimensions["B"].width = 60
    wb.active = 1          # abre na aba Comparação
    return wb, cont


def salvar(wb, pasta_destino, aba):
    carimbo = datetime.datetime.now().strftime("%Y-%m-%d %H%M")
    nome = f"Resultado conferência ISSQN {aba} {carimbo}.xlsx"
    for destino in (pasta_destino, os.path.join(os.path.expanduser("~"), "Desktop"),
                    os.path.expanduser("~")):
        try:
            caminho = os.path.join(destino, nome)
            wb.save(caminho)
            return caminho
        except OSError:
            continue
    raise RuntimeError("Não consegui salvar o resultado (sem permissão de gravação).")


# ============================================================
#  Janelas (tkinter) e configuração lembrada
# ============================================================


# =============================================================================
# Caminhos que funcionam para qualquer usuario do Windows
# =============================================================================
def caminho_do_usuario(texto, base=None):
    """Ajusta um caminho do config para o usuario que esta rodando o script.

    - aceita variaveis do Windows: %USERPROFILE%, %OneDrive%, %USERNAME%, ~
    - caminho relativo (ex.: ..\\Planilhas\\x.xlsx) vale a partir de `base`
    - caminho de OUTRO usuario (C:\\Users\\fulano\\...) que nao existe aqui e
      trocado pela pasta do usuario atual; a parte "OneDrive..." vira a
      OneDrive dele (mesmo que o nome seja "OneDrive - Empresa").
    Assim ninguem precisa editar o config quando outra pessoa usa o script.
    """
    import glob as _glob

    if not texto:
        return texto
    t = str(texto).strip().strip('"')
    t = re.sub(r"%([^%]+)%", lambda m: os.environ.get(m.group(1), m.group(0)), t)
    t = os.path.expanduser(t)
    if base and not os.path.isabs(t) and not t.startswith(("\\\\", "//")):
        t = os.path.normpath(os.path.join(base, t))

    def existe(c):
        return bool(_glob.glob(c)) if any(x in c for x in "*?") else os.path.exists(c)

    if existe(t):
        return t
    m = re.match(r"^[A-Za-z]:[\\/]+Users[\\/]+[^\\/]+[\\/]*(.*)$", t, re.I)
    if not m:
        return t
    resto = m.group(1)
    candidatos = []
    partes = re.split(r"[\\/]+", resto, maxsplit=1)
    for var in ("OneDrive", "OneDriveCommercial", "OneDriveConsumer"):
        od = os.environ.get(var)
        if od and partes[0].lower().startswith("onedrive"):
            candidatos.append(os.path.join(od, partes[1]) if len(partes) > 1 else od)
    candidatos.append(os.path.join(os.path.expanduser("~"), resto))
    for c in candidatos:
        if existe(c):
            return c
    return t


def carregar_cfg():
    try:
        with open(ARQ_CONFIG, encoding="utf-8") as f:
            return json.load(f)
    except (OSError, ValueError):
        return {}


def salvar_cfg(cfg):
    try:
        with open(ARQ_CONFIG, "w", encoding="utf-8") as f:
            json.dump(cfg, f, ensure_ascii=False, indent=2)
    except OSError:
        pass


def tk():
    global _tk_root
    if _tk_root is None:
        import tkinter
        _tk_root = tkinter.Tk()
        _tk_root.withdraw()
        _tk_root.attributes("-topmost", True)
    return _tk_root


def escolher_arquivo(titulo, pasta_inicial):
    from tkinter import filedialog
    return filedialog.askopenfilename(parent=tk(), title=titulo, initialdir=pasta_inicial or None,
                                      filetypes=[("Planilhas Excel", "*.xlsx *.xlsm"), ("Todos", "*.*")])


def escolher_pasta(pasta_inicial):
    from tkinter import filedialog
    return filedialog.askdirectory(parent=tk(), initialdir=pasta_inicial or None,
                                   title="Pasta onde as notas ficam salvas (Cancelar = não verificar)")


def perguntar_aba_gui(abas, sugestao):
    from tkinter import simpledialog
    return simpledialog.askstring("Aba da conferência",
                                  "Qual aba da conferência devo usar?\n\n" + "\n".join(abas),
                                  initialvalue=sugestao or "", parent=tk())


def avisar(msg, erro=False, gui=True, alerta=False):
    print(msg)
    if gui:
        from tkinter import messagebox
        funcao = messagebox.showerror if erro else messagebox.showwarning if alerta else messagebox.showinfo
        funcao("Conferência ISSQN", msg, parent=tk())


def moeda(v):
    return "R$ " + f"{v:,.2f}".replace(",", "X").replace(".", ",").replace("X", ".")


# ============================================================
#  Principal
# ============================================================

def main():
    ap = argparse.ArgumentParser(description="Conferência de retenções ISSQN")
    ap.add_argument("issqn", nargs="?", help="planilha Retenções ISSQN")
    ap.add_argument("conferencia", nargs="?", help="planilha 2-CONFERENCIA RETENÇÕES")
    ap.add_argument("pasta", nargs="?", help="pasta onde as notas ficam salvas")
    ap.add_argument("--aba", help="aba da conferência (ex.: 09_2026)")
    ap.add_argument("--sem-abrir", action="store_true", help="não abrir o resultado no final")
    args = ap.parse_args()

    global PASTA_NOTAS, PASTA_PLANILHAS_ISSQN, PLANILHA_CONFERENCIA
    gui = not args.issqn
    cfg = carregar_cfg()
    # caminhos de outro usuario do Windows viram os deste usuario
    PASTA_NOTAS = caminho_do_usuario(PASTA_NOTAS, PASTA_DO_SCRIPT)
    PASTA_PLANILHAS_ISSQN = caminho_do_usuario(PASTA_PLANILHAS_ISSQN, PASTA_DO_SCRIPT)
    PLANILHA_CONFERENCIA = caminho_do_usuario(PLANILHA_CONFERENCIA, PASTA_DO_SCRIPT)
    for chave in ("pasta_planilhas", "pasta_notas"):
        if cfg.get(chave):
            cfg[chave] = caminho_do_usuario(cfg[chave], PASTA_DO_SCRIPT)
    try:
        # 1) arquivos
        arq_issqn = args.issqn
        arq_conf = args.conferencia
        if not arq_conf and PLANILHA_CONFERENCIA:
            if not os.path.isfile(PLANILHA_CONFERENCIA):
                raise RuntimeError("A planilha de conferência configurada não foi encontrada:\n"
                                   f"{PLANILHA_CONFERENCIA}\n\nConfira PLANILHA_CONFERENCIA no início do script.")
            arq_conf = PLANILHA_CONFERENCIA
        if gui:
            pasta_ini = PASTA_PLANILHAS_ISSQN if os.path.isdir(PASTA_PLANILHAS_ISSQN or "") \
                else cfg.get("pasta_planilhas")
            arq_issqn = escolher_arquivo("Selecione a planilha RETENÇÕES ISSQN", pasta_ini)
            if not arq_issqn:
                return
            cfg["pasta_planilhas"] = os.path.dirname(arq_issqn)
        if not arq_conf:
            if not gui:
                raise RuntimeError("Informe a planilha de conferência (ou preencha PLANILHA_CONFERENCIA).")
            arq_conf = escolher_arquivo("Selecione a planilha 2-CONFERENCIA RETENÇÕES",
                                        os.path.dirname(arq_issqn))
            if not arq_conf:
                return

        # 2) pasta das notas
        pasta = args.pasta
        if not pasta and PASTA_NOTAS:
            if not os.path.isdir(PASTA_NOTAS):
                raise RuntimeError("A pasta das notas configurada não foi encontrada:\n"
                                   f"{PASTA_NOTAS}\n\nConfira PASTA_NOTAS no início do script.")
            pasta = PASTA_NOTAS
        if not pasta and gui:
            pasta = escolher_pasta(cfg.get("pasta_notas"))
            if pasta:
                cfg["pasta_notas"] = pasta
        if pasta and not os.path.isdir(pasta):
            raise RuntimeError(f"A pasta das notas não existe: {pasta}")
        pasta = os.path.normpath(pasta) if pasta else None
        salvar_cfg(cfg)

        # 3) leitura
        print("Lendo planilha ISSQN...")
        notas, aba_issqn, sem_iss_ret = ler_issqn(arq_issqn)
        print("Lendo conferência...")
        conf, aba_conf, sem_q_conf = ler_conferencia(arq_conf, os.path.basename(arq_issqn), args.aba,
                                         perguntar_aba_gui if gui else None)
        if conf is None:
            return
        arquivos, indice = ([], {})
        if pasta:
            print("Lendo arquivos da pasta de notas...")
            arquivos, indice = listar_arquivos(pasta)

        # 4) conferência e resultado
        itens, ignorados = conferir(notas, conf, arquivos, indice, pasta is not None, aba_conf,
                                    sem_iss_ret | sem_q_conf)
        info = {"arq_issqn": os.path.basename(arq_issqn), "aba_issqn": aba_issqn,
                "arq_conf": os.path.basename(arq_conf), "aba_conf": aba_conf,
                "pasta": pasta, "n_arq": len(arquivos)}
        wb, cont = gerar_resultado(itens, conf, info, ignorados)
        caminho = salvar(wb, PASTA_DO_SCRIPT, aba_conf)

        pendentes = [i for i in itens if i["cat"] != "completa"]
        if pendentes:
            linhas_aviso = [f"• {i['nf_txt'] or '(sem nº)'} - {i['nome'] or '(sem nome)'}: {i['falta']}"
                            for i in pendentes[:15]]
            if len(pendentes) > 15:
                linhas_aviso.append(f"... e mais {len(pendentes) - 15}")
            avisar(f"ATENÇÃO: {len(pendentes)} nota(s) com ISSQN com pendência:\n\n"
                   + "\n".join(linhas_aviso)
                   + "\n\nDetalhes na aba \"Comparação\" do resultado.", gui=gui, alerta=True)

        resumo = (f"{cont['total']} notas com ISSQN conferidas.\n\n"
                  f"Completas: {cont['completa']}\n"
                  f"Faltam na Retenções: {cont['f_ret']}\n"
                  f"Faltam na conferência: {cont['f_conf']}\n"
                  f"Faltam na pasta: {cont['f_pasta']}\n"
                  f"Para conferir: {cont['conferir']}\n"
                  f"Valor NÃO bate: {cont['valor']}\n"
                  f"Arquivos na pasta sem nota na Retenções/conferência: {cont['so_pasta']}\n\n"
                  f"Resultado salvo em:\n{caminho}")
        if not args.sem_abrir and hasattr(os, "startfile"):
            os.startfile(caminho)
        avisar(resumo, gui=gui)

    except PermissionError as e:
        avisar(f"O arquivo está bloqueado ou sem permissão:\n{e.filename}\n\n"
               "Se estiver aberto no Excel com alterações, salve e tente de novo.", erro=True, gui=gui)
        sys.exit(1)
    except Exception as e:
        traceback.print_exc()
        avisar(f"Erro: {e}", erro=True, gui=gui)
        sys.exit(1)


if __name__ == "__main__":
    main()
