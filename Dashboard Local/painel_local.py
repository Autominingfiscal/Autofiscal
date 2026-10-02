# -*- coding: utf-8 -*-
"""
PAINEL DA EXPEDIÇÃO DE CONCENTRADO - SOMENTE LOCAL
==================================================
Lê TODAS as abas de dia ("03.09", "04.09", ...) da planilha EXPED_CONCENTRADO e mostra
um painel no navegador DESTE PC.

NÃO USA REDE: não abre porta, não é servidor, não pede liberação de firewall nem de
administrador. Ele só lê a planilha e grava um arquivo HTML nesta pasta:

  - gera "painel_expedicao.html" e abre no navegador;
  - a cada 30 s relê a planilha (só quando ela foi salva) e regrava o arquivo;
  - a página se recarrega sozinha e mostra os dados novos.

Uso:  python painel_local.py        (ou dois cliques em iniciar_painel.bat)
Requisito: openpyxl (já instalado).
"""

import argparse
import configparser
import datetime as dt
import glob
import json
import os
import re
import shutil
import sys
import tempfile
import time
import unicodedata

PASTA = os.path.dirname(os.path.abspath(__file__))
ARQ_CONFIG = os.path.join(PASTA, "painel_config.ini")

CONFIG_PADRAO = r"""; ===== DASHBOARD DA EXPEDIÇÃO - CONFIGURAÇÃO =====
[geral]
; Caminho COMPLETO da planilha. Pode usar * (ex.: ...\EXPED_CONCENTRADO*.xlsx):
; nesse caso o painel usa a mais recente (troca sozinho quando começar um embarque novo).
planilha = C:\CAMINHO\DA\PASTA\EXPED_CONCENTRADO*.xlsx

; De quanto em quanto tempo o painel atualiza (segundos)
atualizar_segundos = 30

; Tema da tela: escuro ou claro
tema = escuro

; Alternar sozinho entre "último dia" e "embarque inteiro" a cada X segundos (0 = não alterna)
rodizio_segundos = 0

; Hora em que começa o turno VESPERTINO
inicio_turno_vespertino = 12:00

; Faixa de aceite sobre o peso bruto alvo (kg)
tolerancia_excesso_kg = 80
tolerancia_subcarga_kg = 1000

[peso_bruto_alvo]
TRUCADO = 48500
CANGURU = 50000
VANDERLEIA = 53000
LS 4 EIXOS = 58500
"""


# =============================================================================
# Configuração
# =============================================================================
def norm(texto):
    if texto is None:
        return ""
    s = unicodedata.normalize("NFKD", str(texto))
    s = "".join(c for c in s if not unicodedata.combining(c))
    return re.sub(r"[^A-Z0-9]", "", s.upper())


def carregar_config():
    if not os.path.exists(ARQ_CONFIG):
        texto = CONFIG_PADRAO
        # aproveita pesos alvo e tolerâncias do lançador, se estiver na mesma pasta
        for nome in ("etapa3_config.ini", "etapa2_config.ini", "etapa1_config.ini"):
            p = os.path.join(PASTA, nome)
            if os.path.exists(p):
                cp = configparser.ConfigParser(interpolation=None)
                try:
                    with open(p, encoding="utf-8-sig") as f:
                        cp.read_file(f)
                except Exception:
                    continue
                g = cp["geral"] if cp.has_section("geral") else {}
                for chave in ("inicio_turno_vespertino", "tolerancia_excesso_kg", "tolerancia_subcarga_kg"):
                    if chave in g:
                        texto = re.sub(rf"(?m)^{chave} = .*$", f"{chave} = {g[chave]}", texto)
                if cp.has_section("peso_bruto_alvo"):
                    texto = texto.split("[peso_bruto_alvo]")[0] + "[peso_bruto_alvo]\n" + "".join(
                        f"{k.upper()} = {v}\n" for k, v in cp["peso_bruto_alvo"].items())
                break
        with open(ARQ_CONFIG, "w", encoding="utf-8") as f:
            f.write(texto)
        print(f"Criei {ARQ_CONFIG}. Coloque nele o caminho da planilha, salve e rode de novo.")
        try:
            os.startfile(ARQ_CONFIG)
        except Exception:
            pass
    cp = configparser.ConfigParser(interpolation=None)
    try:
        with open(ARQ_CONFIG, encoding="utf-8-sig") as f:
            cp.read_file(f)
    except UnicodeDecodeError:
        with open(ARQ_CONFIG, encoding="cp1252") as f:
            cp.read_file(f)
    g = cp["geral"]
    h, m = (g.get("inicio_turno_vespertino", "12:00").strip() + ":0").split(":")[:2]
    alvos = {}
    if cp.has_section("peso_bruto_alvo"):
        for k, v in cp["peso_bruto_alvo"].items():
            try:
                alvos[norm(k)] = float(v)
            except ValueError:
                pass
    return {
        "planilha": g.get("planilha", "").strip().strip('"'),
        "atualizar": max(10, g.getint("atualizar_segundos", 30)),
        "tema": "light" if g.get("tema", "escuro").strip().lower().startswith("clar") else "dark",
        "rodizio": max(0, g.getint("rodizio_segundos", 0)),
        "turno2_min": int(h) * 60 + int(m),
        "tol_exc": g.getfloat("tolerancia_excesso_kg", 80),
        "tol_sub": g.getfloat("tolerancia_subcarga_kg", 1000),
        "alvos": alvos,
    }


# =============================================================================
# Leitura da planilha
# =============================================================================

# =============================================================================
# Caminhos que funcionam para qualquer usuario do Windows
# =============================================================================
def caminho_do_usuario(texto, base=None):
    """Ajusta um caminho do config para o usuario que esta rodando o script.

    - aceita variaveis do Windows: %USERPROFILE%, %OneDrive%, %USERNAME%, ~
    - caminho relativo (ex.: ..\\Planilhas\\x.xlsx) vale a partir de `base`
    - caminho de OUTRO usuario (C:\\Users\\fulano\\...) que nao existe aqui e
      trocado pela pasta do usuario atual; a parte "OneDrive..." vira a
      OneDrive dele (mesmo que o nome seja "OneDrive - Empresa").
    Assim ninguem precisa editar o config quando outra pessoa usa o script.
    """
    import glob as _glob

    if not texto:
        return texto
    t = str(texto).strip().strip('"')
    t = re.sub(r"%([^%]+)%", lambda m: os.environ.get(m.group(1), m.group(0)), t)
    t = os.path.expanduser(t)
    if base and not os.path.isabs(t) and not t.startswith(("\\\\", "//")):
        t = os.path.normpath(os.path.join(base, t))

    def existe(c):
        return bool(_glob.glob(c)) if any(x in c for x in "*?") else os.path.exists(c)

    if existe(t):
        return t
    m = re.match(r"^[A-Za-z]:[\\/]+Users[\\/]+[^\\/]+[\\/]*(.*)$", t, re.I)
    if not m:
        return t
    resto = m.group(1)
    candidatos = []
    partes = re.split(r"[\\/]+", resto, maxsplit=1)
    for var in ("OneDrive", "OneDriveCommercial", "OneDriveConsumer"):
        od = os.environ.get(var)
        if od and partes[0].lower().startswith("onedrive"):
            candidatos.append(os.path.join(od, partes[1]) if len(partes) > 1 else od)
    candidatos.append(os.path.join(os.path.expanduser("~"), resto))
    for c in candidatos:
        if existe(c):
            return c
    return t


def achar_planilha(cfg):
    padrao = caminho_do_usuario(cfg["planilha"], PASTA)
    if any(c in padrao for c in "*?"):
        achados = [p for p in glob.glob(padrao) if not os.path.basename(p).startswith("~$")]
        return max(achados, key=os.path.getmtime) if achados else None
    return padrao if os.path.exists(padrao) else None


def minutos(v):
    """Hora da célula -> minutos desde 00:00 (aceita time, datetime ou número do Excel)."""
    if v is None or v == "":
        return None
    if isinstance(v, dt.datetime):
        v = v.time()
    if isinstance(v, dt.time):
        return v.hour * 60 + v.minute + v.second / 60
    if isinstance(v, dt.timedelta):
        return v.total_seconds() / 60
    try:
        return (float(v) % 1) * 1440
    except (TypeError, ValueError):
        return None


def numero(v):
    try:
        return float(v)
    except (TypeError, ValueError):
        return None


def peso_alvo(cfg, modelo):
    m = norm(modelo)
    if not m:
        return None
    if m in cfg["alvos"]:
        return cfg["alvos"][m]
    for k, v in cfg["alvos"].items():
        if k and (k in m or m in k):
            return v
    return None


def texto(v):
    return "" if v is None else str(v).strip()


CAMPOS = {norm(k): v for k, v in {  # coluna da planilha -> nome no painel
    "TICKET DE PESAGEM": "ticket", "Nº_NOTA": "nf", "CÓDIGO LACRE": "lacre", "TRANSPORTADORA": "transp",
    "MOTORISTA": "motorista", "PLACA DO CAVALO": "placa", "MODELO": "modelo",
    "PESO_ENTRADA": "entrada", "PESO_SAÍDA": "saida", "HORA_ENTRADA": "h_ent", "HORA_SAÍDA": "h_sai",
    "HORA EMISSÃO NF": "h_nf"}.items()}


def ler_planilha(caminho, cfg):
    import openpyxl
    from openpyxl.utils.cell import range_boundaries
    # lê uma CÓPIA (a planilha pode estar aberta no Excel)
    tmp = os.path.join(tempfile.gettempdir(), "dashboard_exped_copia.xlsx")
    shutil.copyfile(caminho, tmp)
    wb = openpyxl.load_workbook(tmp, data_only=True)
    hoje = dt.date.today()
    dias = []
    for ws in wb.worksheets:
        m = re.fullmatch(r"\s*(\d{1,2})\.(\d{1,2})\s*", ws.title)
        if not m:
            continue
        d, mes = int(m.group(1)), int(m.group(2))
        ano = hoje.year - (1 if mes > hoje.month + 1 else 0)
        try:
            data = dt.date(ano, mes, d)
        except ValueError:
            continue
        # tabela de expedição = a que tem a coluna TICKET DE PESAGEM
        ref = None
        for t in ws.tables.values():
            c1, r1, c2, r2 = range_boundaries(t.ref)
            cab = [norm(ws.cell(r1, c).value) for c in range(c1, c2 + 1)]
            if "TICKETDEPESAGEM" in cab:
                ref = (c1, r1, c2, r2, cab)
                break
        if ref is None:
            continue
        c1, r1, c2, r2, cab = ref
        idx = {CAMPOS[h]: i for i, h in enumerate(cab) if h in CAMPOS}
        linhas = []
        for r in range(r1 + 1, r2 + 1):
            vals = [ws.cell(r, c).value for c in range(c1, c2 + 1)]
            get = lambda k: vals[idx[k]] if k in idx else None
            ticket = numero(get("ticket"))
            if not ticket:
                continue
            ent, sai = numero(get("entrada")), numero(get("saida"))
            h_ent, h_sai, h_nf = minutos(get("h_ent")), minutos(get("h_sai")), minutos(get("h_nf"))
            modelo = texto(get("modelo"))
            alvo = peso_alvo(cfg, modelo)
            liq = (sai - ent) if (ent and sai) else None
            if not sai:
                status = "mina"
            elif not alvo:
                status = "semalvo"
            else:
                dif = sai - alvo
                status = ("vermelho" if dif > cfg["tol_exc"] else "amarelo" if dif > 0
                          else "verde" if -dif > cfg["tol_sub"] else "dentro")
            dif_min = lambda a, b: (b - a + 1440) % 1440 if (a is not None and b is not None) else None
            linhas.append({
                "ticket": int(ticket), "nf": int(numero(get("nf"))) if numero(get("nf")) else None,
                "lacre": texto(get("lacre")), "transp": texto(get("transp")), "motorista": texto(get("motorista")),
                "placa": texto(get("placa")), "modelo": modelo or "SEM MODELO",
                "entrada": ent, "saida": sai, "liquido": liq, "alvo": alvo,
                "dif": (sai - alvo) if (sai and alvo) else None,
                "h_ent": h_ent, "h_sai": h_sai, "h_nf": h_nf,
                "galpao": dif_min(h_ent, h_sai), "tempo_nf": dif_min(h_sai, h_nf),
                "turno": 2 if (h_ent is not None and h_ent >= cfg["turno2_min"]) else 1,
                "status": status,
            })
        dias.append({"aba": ws.title.strip(), "data": data.isoformat(), "linhas": linhas})
    dias.sort(key=lambda d: d["data"])
    nome = os.path.basename(caminho)
    emb = re.search(r"(\d+\s*[°ºo]\s*embarque)", nome, re.I)
    return {
        "arquivo": nome,
        "embarque": emb.group(1) if emb else os.path.splitext(nome)[0],
        "salvo_em": dt.datetime.fromtimestamp(os.path.getmtime(caminho)).strftime("%d/%m %H:%M:%S"),
        "lido_em": dt.datetime.now().strftime("%d/%m %H:%M:%S"),
        "dias": dias,
        "tol_exc": cfg["tol_exc"], "tol_sub": cfg["tol_sub"],
        "alvos": {k: v for k, v in cfg["alvos"].items()},
        "atualizar": cfg["atualizar"],
        "tema": cfg["tema"], "rodizio": cfg["rodizio"],
    }


class Dados:
    """Guarda a última leitura; só relê quando a planilha é salva de novo."""

    def __init__(self, cfg):
        self.cfg, self.chave, self.json, self.erro = cfg, None, None, None

    def obter(self):
        if True:
            caminho = achar_planilha(self.cfg)
            if not caminho:
                self.erro = f"Planilha não encontrada: {self.cfg['planilha']}"
                return self.json, self.erro
            chave = (caminho, os.path.getmtime(caminho), os.path.getsize(caminho))
            if chave != self.chave or self.json is None:
                try:
                    self.json = json.dumps(ler_planilha(caminho, self.cfg), ensure_ascii=False)
                    self.chave, self.erro = chave, None
                    print(f"[{dt.datetime.now():%H:%M:%S}] Planilha lida: {os.path.basename(caminho)}", flush=True)
                except Exception as exc:   # ex.: arquivo no meio do salvamento -> usa a leitura anterior
                    self.erro = f"Não consegui ler a planilha agora ({type(exc).__name__}: {exc})"
            return self.json, self.erro


ARQ_PAINEL = os.path.join(PASTA, "painel_expedicao.html")


def gravar_painel(destino, js, erro):
    """Grava o HTML com os dados dentro (troca o arquivo de uma vez, sem deixar pela metade)."""
    dados = json.loads(js) if js else None
    if dados is not None:
        dados["gerado_ts"] = int(time.time() * 1000)
    embutido = json.dumps({"erro": erro, "dados": dados}, ensure_ascii=False).replace("</", "<\\/")
    html = PAGINA.replace("/*DADOS_EMBUTIDOS*/null", embutido)
    tmp = destino + ".tmp"
    with open(tmp, "w", encoding="utf-8") as f:
        f.write(html)
    for _ in range(5):                      # o navegador pode estar lendo o arquivo nesse instante
        try:
            os.replace(tmp, destino)
            return
        except PermissionError:
            time.sleep(0.3)


def painel_local(cfg):
    dados = Dados(cfg)
    print("=" * 70)
    print(" PAINEL DA EXPEDIÇÃO - somente local (sem rede)")
    print(f" Planilha: {cfg['planilha']}")
    print(f" Painel:   {ARQ_PAINEL}")
    print(" O painel abre no navegador e se atualiza sozinho. Tecle F11 para tela cheia.")
    print(" Deixe esta janela aberta (pode minimizar). Ctrl+C para desligar.")
    print("=" * 70, flush=True)
    aberto, ultimo_erro = False, None
    while True:
        try:
            js, erro = dados.obter()
            if erro and erro != ultimo_erro:
                print(f"[{dt.datetime.now():%H:%M:%S}] ATENÇÃO: {erro}", flush=True)
            ultimo_erro = erro
            gravar_painel(ARQ_PAINEL, js, erro)
            if not aberto:
                aberto = True
                try:
                    os.startfile(ARQ_PAINEL)          # abre no navegador padrão (Windows)
                except AttributeError:
                    import webbrowser
                    webbrowser.open("file://" + ARQ_PAINEL)
            time.sleep(cfg["atualizar"])
        except KeyboardInterrupt:
            break
        except Exception as exc:
            print(f"[{dt.datetime.now():%H:%M:%S}] ERRO: {exc}", flush=True)
            time.sleep(cfg["atualizar"])
    print("Painel desligado.")


def main():
    ap = argparse.ArgumentParser(description="Painel da expedição de concentrado (somente local, sem rede)")
    ap.add_argument("--html", metavar="SAIDA.html", help="gera um HTML com os dados de agora e sai")
    args = ap.parse_args()
    cfg = carregar_config()
    if args.html:
        js, erro = Dados(cfg).obter()
        if not js:
            sys.exit(erro)
        gravar_painel(os.path.abspath(args.html), js, erro)
        print(f"Gerado: {args.html}")
    else:
        painel_local(cfg)


# =============================================================================
# Página (HTML + CSS + JS, sem nada da internet)
# =============================================================================
PAGINA = r"""<!doctype html>
<html lang="pt-BR" data-theme="dark">
<head>
<meta charset="utf-8">
<meta name="viewport" content="width=device-width, initial-scale=1">
<title>Expedição de Concentrado</title>
<style>
:root{
  --page:#0d0d0d; --surface:#1a1a19; --ink:#ffffff; --ink2:#c3c2b7; --muted:#898781;
  --grid:#2c2c2a; --axis:#383835; --ring:rgba(255,255,255,.10);
  --s1:#3987e5; --s2:#d95926; --neutral:#6b6a64;
  --good:#0ca30c; --warn:#fab219; --crit:#d03b3b;
}
:root[data-theme="light"]{
  --page:#f9f9f7; --surface:#fcfcfb; --ink:#0b0b0b; --ink2:#52514e; --muted:#898781;
  --grid:#e1e0d9; --axis:#c3c2b7; --ring:rgba(11,11,11,.10);
  --s1:#2a78d6; --s2:#eb6834; --neutral:#a8a79f;
}
*{box-sizing:border-box}
html,body{margin:0;background:var(--page);color:var(--ink);
  font:15px/1.35 system-ui,-apple-system,"Segoe UI",sans-serif}
body{padding:18px 22px;height:100vh;display:flex;flex-direction:column;overflow:hidden}
header{display:flex;align-items:baseline;gap:16px;flex-wrap:wrap;margin-bottom:12px}
h1{font-size:24px;margin:0;font-weight:650}
h1 small{font-weight:400;color:var(--ink2);font-size:18px}
.sp{flex:1}
.estado{color:var(--ink2);font-size:14px;display:flex;align-items:center;gap:8px}
.ponto{width:10px;height:10px;border-radius:50%;background:var(--good)}
.ponto.off{background:var(--crit)}
.filtros{display:flex;gap:8px;align-items:center;flex-wrap:wrap;margin-bottom:14px}
.filtros button,.filtros select{background:var(--surface);color:var(--ink);border:1px solid var(--ring);
  border-radius:8px;padding:7px 12px;font:inherit;cursor:pointer}
.filtros button.sel{border-color:var(--s1);box-shadow:inset 0 0 0 1px var(--s1);font-weight:600}
.aviso{background:#3a1515;color:#ffd7d7;border-radius:8px;padding:8px 12px;margin-bottom:12px;display:none}
:root[data-theme="light"] .aviso{background:#fde8e8;color:#7a1010}
.kpis{display:grid;grid-template-columns:repeat(6,1fr);gap:12px;margin-bottom:12px}
.card{background:var(--surface);border:1px solid var(--ring);border-radius:12px;padding:14px 16px;min-width:0}
.kpi .rot{color:var(--ink2);font-size:14px}
.kpi .val{font-size:40px;font-weight:650;line-height:1.1;margin-top:4px;white-space:nowrap}
.kpi .val span{font-size:18px;font-weight:500;color:var(--ink2);margin-left:4px}
.kpi .sub{color:var(--muted);font-size:13px;margin-top:4px}
.kpi.hero .val{font-size:52px}
.grade{display:grid;grid-template-columns:1.6fr 1fr;gap:12px;margin-bottom:12px}
.grade3{display:grid;grid-template-columns:1.9fr 1fr 1fr;gap:12px;flex:1;min-height:300px}
.grade3 .card{display:flex;flex-direction:column;min-height:0;overflow:hidden}
h2{font-size:16px;margin:0 0 2px;font-weight:600}
.desc{color:var(--muted);font-size:13px;margin:0 0 8px}
.legenda{display:flex;gap:16px;flex-wrap:wrap;color:var(--ink2);font-size:13px;margin:4px 0 6px}
.legenda i{display:inline-block;width:12px;height:12px;border-radius:3px;margin-right:6px;vertical-align:-1px}
svg{display:block;width:100%;overflow:visible}
svg text{fill:var(--muted);font-size:12px;font-variant-numeric:tabular-nums}
svg .val{fill:var(--ink2);font-size:12px;font-weight:600}
.marca{cursor:pointer;transition:opacity .15s}
.marca:hover{opacity:.8}
.faixa-lista{margin-top:10px;display:grid;gap:6px}
.faixa-lista div{display:flex;align-items:center;gap:8px;color:var(--ink2);font-size:14px}
.faixa-lista b{color:var(--ink);font-variant-numeric:tabular-nums;min-width:34px;text-align:right}
.chip{display:inline-flex;align-items:center;gap:5px;border-radius:999px;padding:2px 9px;font-size:12px;
  font-weight:600;white-space:nowrap;color:#111}
.chip svg{width:12px;height:12px}
.c-dentro{background:var(--neutral);color:#fff}.c-verde{background:var(--good);color:#fff}
.c-amarelo{background:var(--warn)}.c-vermelho{background:var(--crit);color:#fff}
.c-mina{background:transparent;color:var(--ink2);border:1px solid var(--axis)}
.c-semalvo{background:transparent;color:var(--muted);border:1px dashed var(--axis)}
table{width:100%;border-collapse:collapse;font-size:14px}
th{color:var(--muted);font-weight:500;text-align:left;padding:4px 6px;border-bottom:1px solid var(--axis);
  font-size:12px;white-space:nowrap}
td{padding:5px 6px;border-bottom:1px solid var(--grid);white-space:nowrap;font-variant-numeric:tabular-nums}
td.n,th.n{text-align:right}
.tabela-wrap{flex:1;min-height:0;overflow:auto}
td.mot{max-width:200px;overflow:hidden;text-overflow:ellipsis}
.barras-h .linha{display:grid;grid-template-columns:190px 1fr 64px;align-items:center;gap:8px;margin:6px 0;font-size:14px}
.barras-h .nome{color:var(--ink2);overflow:hidden;text-overflow:ellipsis;white-space:nowrap}
.barras-h .trilho{height:14px;position:relative}
.barras-h .barra{height:14px;background:var(--s1);border-radius:0 4px 4px 0;min-width:2px}
.barras-h .num{text-align:right;font-variant-numeric:tabular-nums;color:var(--ink)}
#dica{position:fixed;pointer-events:none;background:var(--surface);color:var(--ink);border:1px solid var(--ring);
  border-radius:8px;padding:8px 10px;font-size:13px;box-shadow:0 6px 20px rgba(0,0,0,.35);display:none;z-index:9}
#dica b{font-size:15px}
#dica .k{display:inline-block;width:12px;height:2px;margin-right:6px;vertical-align:middle}
.carregando .card{opacity:.6}
.mina{margin-top:14px;border-top:1px solid var(--grid);padding-top:10px}
.mina h3{font-size:14px;margin:0 0 6px;font-weight:600}
.mina .it{display:flex;gap:10px;font-size:14px;color:var(--ink2);padding:3px 0;font-variant-numeric:tabular-nums}
.mina .it b{color:var(--ink);min-width:86px}
.turnos{display:flex;gap:18px;margin-top:10px;color:var(--ink2);font-size:14px;flex-wrap:wrap}
.turnos i{display:inline-block;width:12px;height:12px;border-radius:3px;margin-right:6px;vertical-align:-1px}
.turnos b{color:var(--ink)}
@media (max-width:1200px){body{height:auto;overflow:visible}.kpis{grid-template-columns:repeat(3,1fr)}.grade,.grade3{grid-template-columns:1fr}}
</style>
</head>
<body>
<header>
  <h1>Expedição de Concentrado <small id="embarque"></small></h1>
  <div class="sp"></div>
  <div class="estado"><span class="ponto" id="ponto"></span><span id="estado">carregando…</span>
    <button id="tema" style="margin-left:8px;background:none;border:1px solid var(--ring);color:var(--ink2);border-radius:6px;padding:3px 8px;cursor:pointer">tema</button></div>
</header>
<div class="aviso" id="aviso"></div>
<div class="filtros" id="filtros">
  <button data-esc="dia" class="sel">Último dia</button>
  <button data-esc="tudo">Embarque inteiro</button>
  <select id="dia" aria-label="Escolher o dia"></select>
</div>

<section class="kpis" id="kpis"></section>

<section class="grade">
  <div class="card">
    <h2>Toneladas embarcadas por dia</h2>
    <p class="desc">Peso líquido (saída − entrada) por turno · passe o mouse para ver os valores</p>
    <div class="legenda"><span><i style="background:var(--s1)"></i>Matutino</span><span><i style="background:var(--s2)"></i>Vespertino</span></div>
    <div id="g-dias"></div>
  </div>
  <div class="card">
    <h2>Faixa de aceite <span class="desc" id="faixa-esc"></span></h2>
    <p class="desc" id="faixa-desc"></p>
    <div id="g-faixa"></div>
    <div class="faixa-lista" id="faixa-lista"></div>
    <div class="turnos" id="turnos"></div>
    <div class="mina"><h3 id="t-mina">Caminhões na mina agora</h3><div id="mina"></div></div>
  </div>
</section>

<section class="grade3">
  <div class="card">
    <h2 id="t-tabela">Pesagens</h2>
    <p class="desc">Mais recentes primeiro</p>
    <div class="tabela-wrap"><table id="tabela"></table></div>
  </div>
  <div class="card barras-h">
    <h2>Por modelo de caminhão</h2>
    <p class="desc" id="desc-modelo">Toneladas embarcadas</p>
    <div id="g-modelo"></div>
  </div>
  <div class="card barras-h">
    <h2>Motoristas com mais viagens</h2>
    <p class="desc" id="desc-mot"></p>
    <div id="g-mot"></div>
  </div>
</section>
<div id="dica"></div>

<script>
const EMBUTIDO = /*DADOS_EMBUTIDOS*/null;
let D = null, esc = "dia", diaSel = null;
// filtro atual guardado no endereço (#esc=tudo ou #esc=dia&dia=18.09) para sobreviver à recarga
const H = new URLSearchParams(location.hash.slice(1));
if(H.get("esc")==="tudo") esc="tudo";
if(H.get("dia")) diaSel=H.get("dia");
function guardarEstado(){ const h=new URLSearchParams(); h.set("esc",esc); if(esc==="dia"&&diaSel&&diaSel!==ultimoDiaComDados()) h.set("dia",diaSel);
  if(document.documentElement.dataset.theme!==((D&&D.tema)||"dark")) h.set("tema",document.documentElement.dataset.theme);
  history.replaceState(null,"","#"+h.toString()); }
const $ = s => document.querySelector(s);
const fmt = (v, c=0) => v==null||isNaN(v) ? "–" : v.toLocaleString("pt-BR",{minimumFractionDigits:c,maximumFractionDigits:c});
const hh = m => m==null ? "–" : String(Math.floor(m/60)%24).padStart(2,"0")+":"+String(Math.floor(m%60)).padStart(2,"0");
const el = (tag, attrs={}, txt) => { const e=document.createElement(tag); for(const k in attrs) e.setAttribute(k,attrs[k]); if(txt!=null) e.textContent=txt; return e; };
const NS = "http://www.w3.org/2000/svg";
const sv = (tag, attrs={}) => { const e=document.createElementNS(NS,tag); for(const k in attrs) e.setAttribute(k,attrs[k]); return e; };
const media = a => { const v=a.filter(x=>x!=null&&x<600); return v.length ? v.reduce((s,x)=>s+x,0)/v.length : null; };

const STATUS = {
  dentro:  {rot:"Dentro da faixa", cor:"var(--neutral)", icone:"M2 6.5l2.5 2.5L10 3"},
  amarelo: {rot:"Excesso até TOLE kg", cor:"var(--warn)", icone:"M6 1.5L11 10.5H1Z"},
  vermelho:{rot:"Excesso acima de TOLE kg", cor:"var(--crit)", icone:"M6 1v6M6 9.5v1"},
  verde:   {rot:"Subcarga (mais de TOLS kg abaixo)", cor:"var(--good)", icone:"M6 1.5v8M2.5 6.5L6 10l3.5-3.5"},
  mina:    {rot:"Na mina (sem saída)", cor:"transparent", icone:"M1 6h10"},
  semalvo: {rot:"Sem peso alvo", cor:"transparent", icone:"M3 3l6 6M9 3l-6 6"},
};
const CURTO = {dentro:"Dentro", amarelo:"Exc. até TOLE", vermelho:"Exc. > TOLE", verde:"Subcarga", mina:"Na mina", semalvo:"Sem alvo"};
const rotulo = (s, curto) => (curto ? CURTO[s] : STATUS[s].rot).replace("TOLE", fmt(D.tol_exc)).replace("TOLS", fmt(D.tol_sub));
const nomeCurto = n => { const p=(n||"").trim().split(/\s+/).filter(x=>!/^(DE|DA|DO|DAS|DOS|E)$/i.test(x));
  return p.length>=3 ? [p[0],p[1],p[p.length-1]].join(" ") : p.join(" "); };
function chip(s, curto){
  const c = el("span",{class:"chip c-"+s});
  const i = sv("svg",{viewBox:"0 0 12 12"}); i.appendChild(sv("path",{d:STATUS[s].icone,stroke:"currentColor","stroke-width":"1.8",fill:"none","stroke-linecap":"round","stroke-linejoin":"round"}));
  c.appendChild(i); c.appendChild(document.createTextNode(rotulo(s, curto))); return c;
}

// ---------- tooltip ----------
const dica = $("#dica");
function mostrarDica(ev, linhas){
  dica.replaceChildren();
  linhas.forEach((l,i)=>{ const d=el("div"); if(l.cor){ const k=el("span",{class:"k"}); k.style.background=l.cor; d.appendChild(k);}
    const b=el(i===0&&!l.cor?"span":"b",{},l.valor); d.appendChild(b); if(l.rot){ d.appendChild(document.createTextNode("  "+l.rot)); } dica.appendChild(d); });
  dica.style.display="block";
  const x = Math.min(ev.clientX+14, innerWidth-dica.offsetWidth-8), y = Math.min(ev.clientY+14, innerHeight-dica.offsetHeight-8);
  dica.style.left=x+"px"; dica.style.top=y+"px";
}
const esconderDica = () => dica.style.display="none";

// ---------- dados do escopo ----------
function linhasEscopo(){
  if(!D) return [];
  if(esc==="tudo") return D.dias.flatMap(d=>d.linhas.map(l=>({...l, aba:d.aba})));
  const d = D.dias.find(d=>d.aba===diaSel) || D.dias[D.dias.length-1];
  return d ? d.linhas.map(l=>({...l, aba:d.aba})) : [];
}
function ultimoDiaComDados(){
  for(let i=D.dias.length-1;i>=0;i--) if(D.dias[i].linhas.length) return D.dias[i].aba;
  return D.dias.length ? D.dias[D.dias.length-1].aba : null;
}

// ---------- KPIs ----------
function kpis(L){
  const feitas = L.filter(l=>l.saida);
  const ton = feitas.reduce((s,l)=>s+(l.liquido||0),0)/1000;
  const comAlvo = feitas.filter(l=>l.status!=="semalvo");
  const dentro = comAlvo.filter(l=>l.status==="dentro").length;
  const mina = L.filter(l=>!l.saida).length, semNf = feitas.filter(l=>!l.nf).length;
  const itens = [
    {rot:"Embarcado", val:fmt(ton,1), un:"t", sub: esc==="tudo" ? `em ${D.dias.filter(d=>d.linhas.length).length} dias` : `dia ${diaSel}`, hero:true},
    {rot:"Viagens", val:fmt(feitas.length), un:"", sub: mina ? `${mina} caminhão(ões) na mina agora` : "nenhum caminhão na mina"},
    {rot:"Média por viagem", val:fmt(feitas.length? ton/feitas.length:null,2), un:"t", sub:"peso líquido"},
    {rot:"Dentro da faixa", val:comAlvo.length? fmt(100*dentro/comAlvo.length,0):"–", un:"%", sub:`${dentro} de ${comAlvo.length} viagens`},
    {rot:"Tempo no galpão", val:fmt(media(feitas.map(l=>l.galpao)),0), un:"min", sub:"média entrada → saída"},
    {rot:"Emissão da NF", val:fmt(media(feitas.map(l=>l.tempo_nf)),0), un:"min", sub: semNf ? `${semNf} viagem(ns) ainda sem NF` : "todas com NF"},
  ];
  const box = $("#kpis"); box.replaceChildren();
  itens.forEach(k=>{ const c=el("div",{class:"card kpi"+(k.hero?" hero":"")}); c.appendChild(el("div",{class:"rot"},k.rot));
    const v=el("div",{class:"val"},k.val); if(k.un){ v.appendChild(el("span",{},k.un)); } c.appendChild(v); c.appendChild(el("div",{class:"sub"},k.sub)); box.appendChild(c); });
}

// ---------- toneladas por dia (colunas empilhadas por turno) ----------
function graficoDias(){
  const box=$("#g-dias"); box.replaceChildren();
  const dias = D.dias.filter(d=>d.linhas.length);
  if(!dias.length){ box.appendChild(el("p",{class:"desc"},"Sem dados.")); return; }
  const W=Math.max(box.clientWidth,500), H=280, m={t:22,r:8,b:28,l:44};
  const serie = dias.map(d=>{ const f=d.linhas.filter(l=>l.saida); const t1=f.filter(l=>l.turno===1).reduce((s,l)=>s+l.liquido,0)/1000, t2=f.filter(l=>l.turno===2).reduce((s,l)=>s+l.liquido,0)/1000; return {aba:d.aba,t1,t2,tot:t1+t2,v1:f.filter(l=>l.turno===1).length,v2:f.filter(l=>l.turno===2).length}; });
  const max = Math.max(...serie.map(s=>s.tot))*1.12 || 1;
  const passo = max>1500?500:max>600?200:max>250?100:50;
  const g = sv("svg",{viewBox:`0 0 ${W} ${H}`,role:"img","aria-label":"Toneladas por dia"});
  const y = v => m.t + (H-m.t-m.b)*(1-v/max);
  for(let v=0; v<=max; v+=passo){ g.appendChild(sv("line",{x1:m.l,x2:W-m.r,y1:y(v),y2:y(v),stroke:v?"var(--grid)":"var(--axis)","stroke-width":1}));
    const t=sv("text",{x:m.l-6,y:y(v)+4,"text-anchor":"end"}); t.textContent=fmt(v); g.appendChild(t); }
  const banda=(W-m.l-m.r)/serie.length, bw=Math.min(24, banda*0.6);
  const destaque = esc==="dia" ? diaSel : null;
  serie.forEach((s,i)=>{
    const cx=m.l+banda*i+banda/2, x=cx-bw/2;
    const apagado = destaque && s.aba!==destaque;
    const grupo=sv("g",{class:"marca",tabindex:0}); if(apagado) grupo.setAttribute("opacity","0.35");
    const y0=y(0), y1=y(s.t1), y2=y(s.t1+s.t2), gap=s.t1&&s.t2?2:0;
    const r=4;
    const seg=(topo,base,cor,arredonda)=>{ const h=base-topo; if(h<=0) return;
      const p = arredonda ? `M${x},${base}V${topo+r}Q${x},${topo} ${x+r},${topo}H${x+bw-r}Q${x+bw},${topo} ${x+bw},${topo+r}V${base}Z` : `M${x},${base}V${topo}H${x+bw}V${base}Z`;
      grupo.appendChild(sv("path",{d:p,fill:cor})); };
    seg(y1, y0, "var(--s1)", !s.t2);
    seg(y2, y1-gap, "var(--s2)", true);
    grupo.appendChild(sv("rect",{x:m.l+banda*i,y:m.t,width:banda,height:H-m.t-m.b,fill:"transparent"}));
    const dicaLinhas=()=>[{valor:`Dia ${s.aba}`},{valor:fmt(s.tot,1)+" t",rot:`total · ${s.v1+s.v2} viagens`},{cor:"var(--s1)",valor:fmt(s.t1,1)+" t",rot:`Matutino (${s.v1})`},{cor:"var(--s2)",valor:fmt(s.t2,1)+" t",rot:`Vespertino (${s.v2})`}];
    grupo.addEventListener("pointermove",ev=>mostrarDica(ev,dicaLinhas())); grupo.addEventListener("pointerleave",esconderDica);
    grupo.addEventListener("click",()=>{ esc="dia"; diaSel=s.aba; render(); });
    g.appendChild(grupo);
    if(serie.length<=16 || s.aba===destaque || i===serie.length-1){ const tv=sv("text",{x:cx,y:y2-6,"text-anchor":"middle",class:"val"}); tv.textContent=fmt(s.tot,0); g.appendChild(tv); }
    if(serie.length<=16 || i%2===0){ const tl=sv("text",{x:cx,y:H-m.b+16,"text-anchor":"middle"}); tl.textContent=s.aba; g.appendChild(tl); }
  });
  box.appendChild(g);
}

// ---------- faixa de aceite ----------
function graficoFaixa(L){
  const ordem=["dentro","amarelo","vermelho","verde"];
  const feitas=L.filter(l=>l.saida && l.status!=="semalvo");
  const cont=Object.fromEntries(ordem.map(s=>[s,feitas.filter(l=>l.status===s).length]));
  const n=feitas.length;
  $("#faixa-esc").textContent = esc==="tudo" ? "· embarque inteiro" : `· dia ${diaSel}`;
  $("#faixa-desc").textContent = `Peso bruto de saída comparado ao alvo do modelo (+${fmt(D.tol_exc)} kg / −${fmt(D.tol_sub)} kg)`;
  const box=$("#g-faixa"); box.replaceChildren();
  const W=Math.max(box.clientWidth,300), H=34;
  const g=sv("svg",{viewBox:`0 0 ${W} ${H}`,role:"img","aria-label":"Faixa de aceite"});
  let x=0; const vivos=ordem.filter(s=>cont[s]);
  vivos.forEach((s,i)=>{ const w=n? (W-2*(vivos.length-1))*cont[s]/n : 0;
    const r=sv("rect",{x,y:4,width:Math.max(w,1),height:26,rx:i===0||i===vivos.length-1?4:0,fill:STATUS[s].cor,class:"marca"});
    r.addEventListener("pointermove",ev=>mostrarDica(ev,[{valor:rotulo(s)},{cor:STATUS[s].cor,valor:`${cont[s]} viagens`,rot:`${fmt(100*cont[s]/n,0)}%`}])); r.addEventListener("pointerleave",esconderDica);
    g.appendChild(r);
    if(w>46){ const t=sv("text",{x:x+w/2,y:22,"text-anchor":"middle"}); t.style.fill = s==="amarelo"?"#111":"#fff"; t.style.fontWeight="600"; t.style.fontSize="13px"; t.textContent=fmt(100*cont[s]/n,0)+"%"; g.appendChild(t); }
    x+=w+2; });
  if(!n){ const t=sv("text",{x:0,y:22}); t.textContent="Nenhuma viagem concluída"; g.appendChild(t); }
  box.appendChild(g);
  const lista=$("#faixa-lista"); lista.replaceChildren();
  ordem.forEach(s=>{ const d=el("div"); d.appendChild(el("b",{},fmt(cont[s]))); d.appendChild(chip(s)); lista.appendChild(d); });
  // matutino x vespertino
  const tb=$("#turnos"); tb.replaceChildren();
  [[1,"Matutino","var(--s1)"],[2,"Vespertino","var(--s2)"]].forEach(([n,nome,cor])=>{ const f=L.filter(l=>l.saida&&l.turno===n);
    const d=el("div"); const i=el("i"); i.style.background=cor; d.appendChild(i); d.appendChild(document.createTextNode(nome+": "));
    d.appendChild(el("b",{},`${f.length} viagens · ${fmt(f.reduce((s,l)=>s+l.liquido,0)/1000,1)} t`)); tb.appendChild(d); });
  // caminhões na mina agora (tara lançada, sem saída) - sempre do último dia
  const ult = D.dias.find(d=>d.aba===ultimoDiaComDados());
  const naMina = ult ? ult.linhas.filter(l=>!l.saida) : [];
  $("#t-mina").textContent = `Caminhões na mina agora (${naMina.length})`;
  const m=$("#mina"); m.replaceChildren();
  if(!naMina.length) m.appendChild(el("div",{class:"it"},"Nenhum caminhão aguardando saída."));
  const agora = new Date(), minAgora = agora.getHours()*60+agora.getMinutes();
  naMina.sort((a,b)=>(a.h_ent??0)-(b.h_ent??0)).slice(0,8).forEach(l=>{ const d=el("div",{class:"it"});
    d.appendChild(el("b",{},l.placa)); d.appendChild(el("span",{},nomeCurto(l.motorista)));
    d.appendChild(el("span",{},`entrou ${hh(l.h_ent)}`)); m.appendChild(d); });
}

// ---------- tabela ----------
function tabela(L){
  $("#t-tabela").textContent = esc==="tudo" ? "Pesagens do embarque" : `Pesagens do dia ${diaSel}`;
  const t=$("#tabela"); t.replaceChildren();
  const comDia = esc==="tudo";
  const cab=[...(comDia?["Dia"]:[]),"Entrada","Saída","Ticket","Placa","Motorista","Modelo","Líquido kg","Dif. alvo","Faixa","NF"];
  const tr=el("tr"); cab.forEach(c=>tr.appendChild(el("th",{class:["Líquido kg","Dif. alvo"].includes(c)?"n":""},c))); t.appendChild(el("thead")).appendChild(tr);
  const tb=el("tbody");
  const ord=[...L];
  const chave=l=>(l.aba.split(".").reverse().join(""))+String(Math.round((l.h_sai??l.h_ent??0))).padStart(5,"0");
  ord.sort((a,b)=>chave(b).localeCompare(chave(a)));
  ord.slice(0,esc==="tudo"?200:80).forEach(l=>{ const r=el("tr");
    [...(comDia?[l.aba]:[]),hh(l.h_ent),hh(l.h_sai),l.ticket,l.placa].forEach(v=>r.appendChild(el("td",{},v)));
    const tm=el("td",{class:"mot",title:l.motorista},nomeCurto(l.motorista)); r.appendChild(tm); r.appendChild(el("td",{},l.modelo));
    r.appendChild(el("td",{class:"n"},fmt(l.liquido)));
    r.appendChild(el("td",{class:"n"},l.dif==null?"–":(l.dif>0?"+":"")+fmt(l.dif)));
    const c=el("td"); c.appendChild(chip(l.status, true)); r.appendChild(c);
    r.appendChild(el("td",{},l.nf??"–")); tb.appendChild(r); });
  t.appendChild(tb);
}

// ---------- barras horizontais ----------
function barras(box, itens, unidade, casas){
  box.replaceChildren();
  const max=Math.max(...itens.map(i=>i.v),0)||1;
  itens.forEach(i=>{ const l=el("div",{class:"linha"}); l.appendChild(el("div",{class:"nome",title:i.nome},i.nome));
    const tr=el("div",{class:"trilho"}); const b=el("div",{class:"barra marca",tabindex:0}); b.style.width=(100*i.v/max)+"%";
    b.addEventListener("pointermove",ev=>mostrarDica(ev,[{valor:i.nome},{cor:"var(--s1)",valor:fmt(i.v,casas)+" "+unidade,rot:i.extra||""}])); b.addEventListener("pointerleave",esconderDica);
    tr.appendChild(b); l.appendChild(tr); l.appendChild(el("div",{class:"num"},fmt(i.v,casas))); box.appendChild(l); });
  if(!itens.length) box.appendChild(el("p",{class:"desc"},"Sem dados."));
}
function porModelo(L){
  const f=L.filter(l=>l.saida), grupos={};
  f.forEach(l=>{ const g=grupos[l.modelo]||(grupos[l.modelo]={nome:l.modelo,v:0,n:0}); g.v+=l.liquido/1000; g.n++; });
  const itens=Object.values(grupos).sort((a,b)=>b.v-a.v).map(g=>({...g,extra:`${g.n} viagens · média ${fmt(g.v/g.n,2)} t`}));
  $("#desc-modelo").textContent="Toneladas embarcadas "+(esc==="tudo"?"no embarque":`no dia ${diaSel}`);
  barras($("#g-modelo"), itens, "t", 1);
}
function motoristas(L){
  const f=L.filter(l=>l.saida), grupos={};
  f.forEach(l=>{ const k=l.motorista||l.placa; const g=grupos[k]||(grupos[k]={nome:nomeCurto(k),v:0,t:0}); g.v++; g.t+=l.liquido/1000; });
  const itens=Object.values(grupos).sort((a,b)=>b.v-a.v||b.t-a.t).slice(0,10).map(g=>({...g,extra:`${fmt(g.t,1)} t`}));
  $("#desc-mot").textContent="Viagens concluídas "+(esc==="tudo"?"no embarque":`no dia ${diaSel}`);
  barras($("#g-mot"), itens, "viagens", 0);
}

// ---------- tela ----------
function render(){
  if(!D) return;
  $("#embarque").textContent = "· " + D.embarque;
  const sel=$("#dia"); const atual=sel.value; sel.replaceChildren();
  [...D.dias].reverse().forEach(d=>{ const o=el("option",{value:d.aba},`Dia ${d.aba} (${d.linhas.filter(l=>l.saida).length} viagens)`); sel.appendChild(o); });
  if(!diaSel || !D.dias.some(d=>d.aba===diaSel)) diaSel = ultimoDiaComDados();
  sel.value = diaSel;
  document.querySelectorAll("#filtros button").forEach(b=>b.classList.toggle("sel", b.dataset.esc===esc && !(esc==="dia" && diaSel!==ultimoDiaComDados())));
  const L=linhasEscopo();
  kpis(L); graficoDias(); graficoFaixa(L); tabela(L); porModelo(L); motoristas(L);
  guardarEstado();
}

async function atualizar(){
  document.body.classList.add("carregando");
  try{
    const r = EMBUTIDO || {erro:"Abra o painel pelo programa painel_local.py", dados:null};
    const aviso=$("#aviso");
    if(r.erro){ aviso.textContent=r.erro; aviso.style.display="block"; } else aviso.style.display="none";
    if(r.dados){ const acompanhaUltimo = esc==="dia" && D && diaSel===ultimoDiaComDados(); D=r.dados;
      if(!H.get("tema") && !D._temaAplicado){ document.documentElement.dataset.theme = D.tema || "dark"; D._temaAplicado=true; }
      if(acompanhaUltimo) diaSel=ultimoDiaComDados(); render();
      const idade = D.gerado_ts ? (Date.now()-D.gerado_ts)/1000 : 0;
      const parado = EMBUTIDO && idade > 3*(D.atualizar||30);
      $("#estado").textContent = parado
        ? `o programa do painel parou? dados de ${new Date(D.gerado_ts).toLocaleTimeString("pt-BR")} (planilha salva às ${D.salvo_em})`
        : `planilha salva às ${D.salvo_em} · atualizado ${new Date(D.gerado_ts||Date.now()).toLocaleTimeString("pt-BR")}`;
      $("#ponto").classList.toggle("off", parado); }
  }catch(e){ $("#estado").textContent="não consegui carregar os dados"; $("#ponto").classList.add("off"); }
  document.body.classList.remove("carregando");
  const seg = ((D&&D.atualizar)||30)*1000;
  setTimeout(()=>location.reload(), seg);   // recarrega o arquivo que o programa regrava
}
document.querySelectorAll("#filtros button").forEach(b=>b.addEventListener("click",()=>{ esc=b.dataset.esc; if(esc==="dia") diaSel=ultimoDiaComDados(); render(); }));
$("#dia").addEventListener("change",e=>{ esc="dia"; diaSel=e.target.value; render(); });
$("#tema").addEventListener("click",()=>{ const r=document.documentElement; r.dataset.theme = r.dataset.theme==="dark"?"light":"dark"; render(); });
const Q = new URLSearchParams(location.search);
if(Q.get("tema")==="claro"||H.get("tema")==="light") document.documentElement.dataset.theme="light";
if(H.get("tema")==="dark") document.documentElement.dataset.theme="dark";
addEventListener("resize",()=>render());
atualizar().then(()=>{
  const rod = parseInt(Q.get("rodizio") || (D&&D.rodizio) || "0");
  if(rod>=5) setInterval(()=>{ esc = esc==="dia" ? "tudo" : "dia"; if(esc==="dia") diaSel=ultimoDiaComDados(); render(); }, rod*1000);
});
</script>
</body>
</html>
"""

if __name__ == "__main__":
    main()
