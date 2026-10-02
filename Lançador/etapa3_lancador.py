# -*- coding: utf-8 -*-
"""
LANÇADOR DA BALANÇA - ETAPAS 1, 2 E 3
=====================================
Fica rodando com a planilha EXPED_CONCENTRADO ABERTA no Excel. A cada X segundos
olha as pastas dos tickets e atualiza a aba do dia.

ETAPA 1 - TICKET DE TARA (caminhão entrando vazio; peso no campo "Líquido")
  Tabela de expedição, colunas E até N e S:
    E TICKET | F LACRE | G TRANSPORTADORA | H MOTORISTA | I PLACA CAVALO | J UF |
    K PLACA REBOQUE | L UF | M MODELO  -> da tabela do turno (matutino/vespertino)
    N PESO_ENTRADA | S HORA_ENTRADA    -> do ticket de tara
  Tabela do turno, coluna PESO A CARREGAR = peso bruto alvo do modelo - peso de entrada

ETAPA 2 - TICKET COMPLETO (caminhão voltando carregado)
  Na linha do mesmo ticket na expedição:
    O PESO_SAÍDA (bruto) | T HORA_SAÍDA
  e confere o peso bruto com o alvo do modelo (tolerância +80 kg / -1.000 kg),
  pintando a célula do PESO_SAÍDA:
    sem cor  = dentro do limite (de 1.000 kg abaixo do alvo até o alvo)
    VERDE    = subcarregado (mais de 1.000 kg abaixo do alvo)
    AMARELO  = excesso de 1 a 80 kg
    VERMELHO = excesso acima de 80 kg

ETAPA 3 - NOTA FISCAL (DANFE em PDF)
  Na linha do caminhão na expedição:
    D Nº_NOTA | V HORA EMISSÃO NF
  A nota é ligada à linha pelo nº do TICKET escrito na nota; se não tiver, pela PLACA
  (linha já com peso de saída e ainda sem nota) e, se houver empate, pelo PESO.
  Nº da nota: tirado da chave de acesso (44 dígitos).
  Hora (coluna V): SOMENTE do campo "PROTOCOLO DE AUTORIZAÇÃO DE USO". Sem protocolo, a hora fica em branco.

Regras gerais:
  - Aba: a do dia da ENTRADA do ticket (ex.: 25/09 -> aba "25.09"). Nenhuma aba é criada.
  - Turno: 1ª tabela abaixo da expedição = MATUTINO, 2ª = VESPERTINO (entrada antes de 12:00 = matutino).
  - Caminhão fora das tabelas de turno: avisa e espera você incluir.
  - Não mexe nas fórmulas nem nas outras colunas.

Uso:
    python etapa3_lancador.py                      -> fica monitorando (Ctrl+C para parar)
    python etapa3_lancador.py --teste ARQ.xlsx     -> só mostra o que faria (não grava nada)
    python etapa3_lancador.py --ver-nota NOTA.pdf  -> mostra o que ele leu de uma nota

Requisitos: pywin32 e pypdf  (pip install pywin32 pypdf)
"""

import argparse
import configparser
import csv
import datetime as dt
import fnmatch
import os
import re
import sys
import time
import traceback
import unicodedata

PASTA_SCRIPT = os.path.dirname(os.path.abspath(__file__))
ARQ_CONFIG = os.path.join(PASTA_SCRIPT, "etapa3_config.ini")
ARQ_CONFIG_ETAPA2 = os.path.join(PASTA_SCRIPT, "etapa2_config.ini")
ARQ_CONFIG_ETAPA1 = os.path.join(PASTA_SCRIPT, "etapa1_config.ini")
ARQ_LOG = os.path.join(PASTA_SCRIPT, "etapa3_log.csv")
ARQ_ERROS = os.path.join(PASTA_SCRIPT, "etapa3_erros.log")

CONFIG_PADRAO = """\
; ===== LANÇADOR (ETAPAS 1, 2 E 3) - CONFIGURAÇÃO =====
; Edite no Bloco de Notas, salve e reinicie o programa.

[geral]
; Pasta onde a balança salva os tickets de TARA (caminhão entrando vazio)
pasta_tickets_tara = {pasta}

; Pasta onde a balança salva os tickets COMPLETOS (caminhão voltando carregado)
pasta_tickets_completos = {pasta_completos}

; Pasta onde ficam as NOTAS FISCAIS (DANFE em PDF)
pasta_notas_fiscais = {pasta_notas}

; Nome do arquivo da planilha ABERTA no Excel (pode usar * no fim)
planilha = EXPED_CONCENTRADO*

; De quanto em quanto tempo olhar a pasta (segundos)
intervalo_segundos = 15

; Hora em que começa o turno VESPERTINO (antes disso = MATUTINO)
inicio_turno_vespertino = 12:00

; Só lança tickets cujo Produto contenha este texto (deixe vazio para não filtrar)
produto_contem = CONCENTRADO

; Salvar a planilha depois de lançar
salvar_apos_lancar = sim

; Tolerâncias sobre o peso bruto alvo (kg)
tolerancia_excesso_kg = 80
tolerancia_subcarga_kg = 1000

; Apitar quando passar do excesso ou da subcarga permitidos
alerta_sonoro = sim

[peso_bruto_alvo]
; Peso bruto alvo por MODELO (o nome tem que bater com a coluna MODELO das tabelas de turno)
TRUCADO = 48500
CANGURU = 50000
VANDERLEIA = 53000
LS 4 EIXOS = 58500
"""

ALVOS_PADRAO = {"TRUCADO": 48500, "CANGURU": 50000, "VANDERLEIA": 53000, "LS 4 EIXOS": 58500}


# =============================================================================
# Utilidades
# =============================================================================
def norm(texto):
    """Maiúsculas, sem acento, só letras e números: 'PLACA DO \\nCAVALO' -> 'PLACADOCAVALO'."""
    if texto is None:
        return ""
    s = unicodedata.normalize("NFKD", str(texto))
    s = "".join(c for c in s if not unicodedata.combining(c))
    return re.sub(r"[^A-Z0-9]", "", s.upper())


def formatar_placa(placa):
    p = norm(placa)
    return f"{p[:3]}-{p[3:]}" if len(p) == 7 else (placa or "").strip()


def vazio(v):
    return v is None or (isinstance(v, str) and v.strip() == "")


def como_int(v):
    try:
        return int(float(v))
    except (TypeError, ValueError):
        return None


def nome_aba(data):
    return f"{data.day:02d}.{data.month:02d}"


def kg(v):
    return f"{v:,.0f}".replace(",", ".")


_ja_mostrado = {}


def mostrar(msg, chave=None, repetir_apos=300):
    """Imprime com hora. Com 'chave', não repete a mesma mensagem antes de 'repetir_apos' segundos."""
    if chave:
        if time.time() - _ja_mostrado.get(chave, 0) < repetir_apos:
            return
        _ja_mostrado[chave] = time.time()
    print(f"[{dt.datetime.now():%H:%M:%S}] {msg}", flush=True)



# =============================================================================
# Caminhos que funcionam para qualquer usuario do Windows
# =============================================================================
def caminho_do_usuario(texto, base=None):
    """Ajusta um caminho do config para o usuario que esta rodando o script.

    - aceita variaveis do Windows: %USERPROFILE%, %OneDrive%, %USERNAME%, ~
    - caminho relativo (ex.: ..\\Planilhas\\x.xlsx) vale a partir de `base`
    - caminho de OUTRO usuario (C:\\Users\\fulano\\...) que nao existe aqui e
      trocado pela pasta do usuario atual; a parte "OneDrive..." vira a
      OneDrive dele (mesmo que o nome seja "OneDrive - Empresa").
    Assim ninguem precisa editar o config quando outra pessoa usa o script.
    """
    import glob as _glob

    if not texto:
        return texto
    t = str(texto).strip().strip('"')
    t = re.sub(r"%([^%]+)%", lambda m: os.environ.get(m.group(1), m.group(0)), t)
    t = os.path.expanduser(t)
    if base and not os.path.isabs(t) and not t.startswith(("\\\\", "//")):
        t = os.path.normpath(os.path.join(base, t))

    def existe(c):
        return bool(_glob.glob(c)) if any(x in c for x in "*?") else os.path.exists(c)

    if existe(t):
        return t
    m = re.match(r"^[A-Za-z]:[\\/]+Users[\\/]+[^\\/]+[\\/]*(.*)$", t, re.I)
    if not m:
        return t
    resto = m.group(1)
    candidatos = []
    partes = re.split(r"[\\/]+", resto, maxsplit=1)
    for var in ("OneDrive", "OneDriveCommercial", "OneDriveConsumer"):
        od = os.environ.get(var)
        if od and partes[0].lower().startswith("onedrive"):
            candidatos.append(os.path.join(od, partes[1]) if len(partes) > 1 else od)
    candidatos.append(os.path.join(os.path.expanduser("~"), resto))
    for c in candidatos:
        if existe(c):
            return c
    return t


def carregar_config():
    if not os.path.exists(ARQ_CONFIG) and os.path.exists(ARQ_CONFIG_ETAPA2):
        # aproveita TODO o ini da etapa 2 e só acrescenta a pasta das notas
        try:
            with open(ARQ_CONFIG_ETAPA2, encoding="utf-8-sig") as f:
                texto = f.read()
        except UnicodeDecodeError:
            with open(ARQ_CONFIG_ETAPA2, encoding="cp1252") as f:
                texto = f.read()
        bloco = ("\n; Pasta onde ficam as NOTAS FISCAIS (DANFE em PDF)\n"
                 "pasta_notas_fiscais = C:\\CAMINHO\\DA\\PASTA\\DAS\\NOTAS_FISCAIS\n")
        m = re.search(r"(?m)^pasta_tickets_completos\s*=.*$", texto)
        texto = texto[:m.end()] + "\n" + bloco + texto[m.end():] if m else texto.replace("[geral]", "[geral]" + bloco, 1)
        with open(ARQ_CONFIG, "w", encoding="utf-8") as f:
            f.write(texto)
        mostrar(f"Criei {os.path.basename(ARQ_CONFIG)} a partir da etapa 2. "
                f"Coloque nele a pasta das NOTAS FISCAIS, salve e rode de novo.")
        try:
            os.startfile(ARQ_CONFIG)
        except Exception:
            pass
    if not os.path.exists(ARQ_CONFIG):
        # aproveita o que já foi configurado na etapa 1 (pasta de tara, planilha, turno, alvos)
        anterior = configparser.ConfigParser(interpolation=None)
        if os.path.exists(ARQ_CONFIG_ETAPA1):
            try:
                with open(ARQ_CONFIG_ETAPA1, encoding="utf-8-sig") as f:
                    anterior.read_file(f)
            except (UnicodeDecodeError, configparser.Error):
                anterior = configparser.ConfigParser(interpolation=None)
        ga = anterior["geral"] if anterior.has_section("geral") else {}
        texto = CONFIG_PADRAO.format(pasta=ga.get("pasta_tickets_tara", PASTA_SCRIPT),
                                     pasta_completos=r"C:\CAMINHO\DA\PASTA\DOS\TICKETS_COMPLETOS",
                                     pasta_notas=r"C:\\CAMINHO\\DA\\PASTA\\DAS\\NOTAS_FISCAIS")
        for chave in ("planilha", "intervalo_segundos", "inicio_turno_vespertino", "produto_contem",
                      "salvar_apos_lancar"):
            if chave in ga:
                texto = re.sub(rf"(?m)^{chave} = .*$", lambda _m, c=chave: f"{c} = {ga[c]}", texto)
        if anterior.has_section("peso_bruto_alvo"):
            texto = texto.split("[peso_bruto_alvo]")[0] + "[peso_bruto_alvo]\n" + "".join(
                f"{k.upper()} = {v}\n" for k, v in anterior["peso_bruto_alvo"].items())
        with open(ARQ_CONFIG, "w", encoding="utf-8") as f:
            f.write(texto)
        mostrar(f"Criei {os.path.basename(ARQ_CONFIG)} (com o que já estava na etapa 1). "
                f"Coloque nele a pasta dos tickets COMPLETOS, salve e rode de novo.")
        try:
            os.startfile(ARQ_CONFIG)       # abre no Bloco de Notas (Windows)
        except Exception:
            pass
    cp = configparser.ConfigParser(interpolation=None)
    try:
        with open(ARQ_CONFIG, encoding="utf-8-sig") as f:
            cp.read_file(f)
    except UnicodeDecodeError:
        with open(ARQ_CONFIG, encoding="cp1252") as f:
            cp.read_file(f)
    g = cp["geral"]
    hora = g.get("inicio_turno_vespertino", g.get("inicio_turno_2", "12:00"))
    h, m = (hora.strip() + ":0").split(":")[:2]
    alvos = dict(ALVOS_PADRAO)
    if cp.has_section("peso_bruto_alvo"):
        alvos = {}
        for modelo, valor in cp["peso_bruto_alvo"].items():
            try:
                alvos[modelo.upper()] = float(valor)
            except ValueError:
                pass
    return {
        "alvos": alvos,
        "pasta": caminho_do_usuario(g.get("pasta_tickets_tara", ""), PASTA_SCRIPT),
        "pasta_completos": caminho_do_usuario(g.get("pasta_tickets_completos", ""), PASTA_SCRIPT),
        "pasta_notas": caminho_do_usuario(g.get("pasta_notas_fiscais", ""), PASTA_SCRIPT),
        "tol_exc": g.getfloat("tolerancia_excesso_kg", 80),
        "tol_sub": g.getfloat("tolerancia_subcarga_kg", 1000),
        "bipe": g.get("alerta_sonoro", "sim").strip().lower() in ("sim", "s", "1", "true"),
        "planilha": g.get("planilha", "EXPED_CONCENTRADO*").strip().strip('"'),
        "intervalo": max(5, g.getint("intervalo_segundos", 15)),
        "turno2": dt.time(int(h), int(m)),
        "produto": g.get("produto_contem", "").strip(),
        "salvar": g.get("salvar_apos_lancar", "sim").strip().lower() in ("sim", "s", "1", "true"),
    }


def gravar_log(resultado, arquivo, ticket="", aba="", linha="", obs=""):
    novo = not os.path.exists(ARQ_LOG)
    try:
        with open(ARQ_LOG, "a", encoding="utf-8-sig", newline="") as f:
            w = csv.writer(f, delimiter=";")
            if novo:
                w.writerow(["data_hora", "resultado", "arquivo", "ticket", "aba", "linha", "observacao"])
            w.writerow([f"{dt.datetime.now():%d/%m/%Y %H:%M:%S}", resultado, arquivo, ticket, aba, linha, obs])
    except OSError:
        pass  # log aberto no Excel, por exemplo: não trava o programa


# =============================================================================
# Leitura do PDF do ticket (FortesReport: texto sem compressão)
# =============================================================================
_RE_TEXTO = re.compile(
    rb"/(F\d+)\s+([\d.]+)\s+Tf\s+(-?[\d.]+)\s+(-?[\d.]+)\s+Td\s*\(((?:\\.|[^\\)])*)\)\s*Tj", re.S)
_ESC = {b"n": b"\n", b"r": b"\r", b"t": b"\t", b"b": b"\b", b"f": b"\f"}


def _desescapar(raw):
    out, i = bytearray(), 0
    while i < len(raw):
        c = raw[i:i + 1]
        if c == b"\\" and i + 1 < len(raw):
            n = raw[i + 1:i + 2]
            m = re.match(rb"[0-7]{1,3}", raw[i + 1:])
            if m:                               # \ddd octal ("\8" e "\9" nao sao octal)
                out.append(int(m.group(0), 8) & 0xFF)
                i += 1 + len(m.group(0))
                continue
            out += _ESC.get(n, n)
            i += 2
            continue
        out += c
        i += 1
    return out.decode("cp1252", errors="replace")


def _textos(dados):
    return [{"fonte": m.group(1).decode(), "x": float(m.group(3)), "y": float(m.group(4)),
             "texto": _desescapar(m.group(5)).strip()} for m in _RE_TEXTO.finditer(dados)]


def _valor(itens, rotulo):
    """Texto impresso à direita do rótulo, na mesma linha."""
    alvo = norm(rotulo)
    for r in itens:
        if norm(r["texto"]) != alvo:
            continue
        cands = [v for v in itens if v is not r and abs(v["y"] - r["y"]) < 2
                 and r["x"] < v["x"] < r["x"] + 165 and v["fonte"] != r["fonte"]]
        if cands:
            return min(cands, key=lambda v: v["x"])["texto"]
    return ""


def _assinatura(itens, rotulo):
    """Nome impresso logo acima da linha de assinatura."""
    alvo = norm(rotulo)
    for r in itens:
        if norm(r["texto"]) == alvo:
            for v in itens:
                if v is not r and 0 < v["y"] - r["y"] < 15 and abs(v["x"] - r["x"]) < 120 and v["texto"]:
                    return v["texto"]
    return ""


def _peso(txt):
    s = re.sub(r"[^0-9,.-]", "", txt or "").replace(".", "").replace(",", ".")
    try:
        return float(s)
    except ValueError:
        return 0.0


def _data_hora(txt):
    m = re.search(r"(\d{2})/(\d{2})/(\d{2,4})\s+(\d{1,2}):(\d{2})", txt or "")
    if not m:
        return None
    d, mes, a, h, mi = m.groups()
    try:
        return dt.datetime(int(a) + (2000 if len(a) == 2 else 0), int(mes), int(d), int(h), int(mi))
    except ValueError:                          # data impossivel (31/02, 25:00...)
        return None


def ler_ticket(caminho):
    """Dados do ticket de tara, ou None se não deu para ler."""
    with open(caminho, "rb") as f:
        itens = _textos(f.read())
    if not itens:
        return None
    tara = _peso(_valor(itens, "Tara"))
    liquido = _peso(_valor(itens, "Líquido"))
    bruto = _peso(_valor(itens, "Bruto"))
    t = {
        "bruto": bruto,
        "tara": tara,
        "saida": _data_hora(_valor(itens, "Data Saída")),
        "numero": como_int(_valor(itens, "NºTicket")),
        "produto": _valor(itens, "Produto"),
        "transportadora": _valor(itens, "Transportadora"),
        "cavalo": _valor(itens, "Cavalo"),
        "reboque": _valor(itens, "1º Reboque"),
        "entrada": _data_hora(_valor(itens, "Data Entrada")),
        "motorista": _assinatura(itens, "Assinatura do Motorista"),
        # ticket de tara: o peso vem no campo "Líquido" e a Tara fica vazia.
        # (se cair aqui um ticket completo, vale o campo Tara)
        "peso_entrada": tara or liquido or bruto,
    }
    if not t["numero"] or not t["cavalo"] or not t["entrada"] or t["peso_entrada"] <= 0:
        return None
    # COMPLETO = tem tara e bruto (bruto maior que a tara) e data de saída; senão é ticket de TARA
    t["tipo"] = "completo" if (tara > 0 and bruto > tara and t["saida"]) else "tara"
    return t


# =============================================================================
# Planilha
# =============================================================================
class Tabela:
    def __init__(self, nome, linha_cab, col_ini, cab, linhas):
        self.nome, self.linha_cab, self.col_ini = nome, linha_cab, col_ini
        self.cab = [norm(c) for c in cab]
        self.linhas = linhas            # valores das linhas de dados

    def col(self, nome):
        n = norm(nome)
        return self.cab.index(n) if n in self.cab else None

    def linha_excel(self, i):
        return self.linha_cab + 1 + i


class PlanilhaExcel:
    """A planilha aberta no Excel (via pywin32). Lê cada aba uma vez por ciclo."""

    def __init__(self, wb):
        self.wb, self.nome = wb, wb.Name
        self.pendente = False           # gravou algo que ainda nao foi salvo
        self.limpar_cache()

    def limpar_cache(self):
        self._abas, self._tabs = None, {}

    def viva(self):
        try:
            return bool(self.wb.Name)
        except Exception:
            return False

    def abas(self):
        if self._abas is None:
            self._abas = [s.Name for s in self.wb.Worksheets]
        return self._abas

    def tabelas(self, aba):
        if aba not in self._tabs:
            res = []
            for lo in self.wb.Worksheets(aba).ListObjects:
                cab = list(lo.HeaderRowRange.Value[0])
                corpo, linhas = lo.DataBodyRange, []
                if corpo is not None:
                    v = corpo.Value
                    linhas = [list(r) for r in (v if isinstance(v, tuple) else ((v,),))]
                res.append(Tabela(lo.Name, lo.Range.Row, lo.Range.Column, cab, linhas))
            self._tabs[aba] = res
        return self._tabs[aba]

    def ler(self, aba, linha, col):
        return self.wb.Worksheets(aba).Cells(linha, col).Value

    def escrever(self, aba, linha, col, valores):
        sh = self.wb.Worksheets(aba)
        sh.Range(sh.Cells(linha, col), sh.Cells(linha, col + len(valores) - 1)).Value = (tuple(valores),)
        self._tabs.pop(aba, None)
        self.pendente = True

    def adicionar_linha(self, aba, nome_tabela):
        self._tabs.pop(aba, None)
        self.pendente = True
        return self.wb.Worksheets(aba).ListObjects(nome_tabela).ListRows.Add().Range.Row

    def pintar(self, aba, linha, col, cor):
        """Cor de fundo e da fonte de uma célula (cor = (fundo_RGB, fonte_RGB, negrito))."""
        fundo, fonte, negrito = cor
        rgb = lambda h: int(h[0:2], 16) + int(h[2:4], 16) * 256 + int(h[4:6], 16) * 65536
        c = self.wb.Worksheets(aba).Cells(linha, col)
        c.Interior.Color = rgb(fundo)
        c.Font.Color = rgb(fonte)
        c.Font.Bold = negrito
        self.pendente = True

    def tirar_cor(self, aba, linha, col):
        """Volta a célula ao visual da tabela: sem preenchimento e com a fonte igual à da célula ao lado."""
        c = self.wb.Worksheets(aba).Cells(linha, col)
        if c.Interior.ColorIndex == -4142:     # já está sem cor
            return
        c.Interior.ColorIndex = -4142           # xlNone: a cor da faixa da tabela volta a aparecer
        self.pendente = True
        try:
            vizinha = self.wb.Worksheets(aba).Cells(linha, col - 1).DisplayFormat.Font
            c.Font.Color, c.Font.Bold = vizinha.Color, vizinha.Bold
        except Exception:
            c.Font.ColorIndex = -4105           # automática

    def salvar(self):
        self.wb.Save()
        self.pendente = False


class PlanilhaTeste:
    """Modo --teste: lê o arquivo salvo e só MOSTRA o que seria gravado."""

    def __init__(self, caminho):
        import openpyxl
        self.wb, self.nome = openpyxl.load_workbook(caminho), os.path.basename(caminho)

    def limpar_cache(self):
        pass

    def viva(self):
        return True

    def abas(self):
        return self.wb.sheetnames

    def tabelas(self, aba):
        from openpyxl.utils.cell import range_boundaries
        ws, res = self.wb[aba], []
        for t in ws.tables.values():
            c1, r1, c2, r2 = range_boundaries(t.ref)
            res.append(Tabela(t.name, r1, c1, [ws.cell(r1, c).value for c in range(c1, c2 + 1)],
                              [[ws.cell(r, c).value for c in range(c1, c2 + 1)] for r in range(r1 + 1, r2 + 1)]))
        return res

    def ler(self, aba, linha, col):
        return self.wb[aba].cell(linha, col).value

    def escrever(self, aba, linha, col, valores):
        from openpyxl.utils import get_column_letter
        for k, v in enumerate(valores):
            self.wb[aba].cell(linha, col + k).value = v
        print(f"      [teste] {aba}!{get_column_letter(col)}{linha}:"
              f"{get_column_letter(col + len(valores) - 1)}{linha} = {valores}")

    def adicionar_linha(self, aba, nome_tabela):
        raise RuntimeError("tabela de expedição cheia (no modo normal uma linha seria adicionada)")

    def pintar(self, aba, linha, col, cor):
        from openpyxl.styles import Font, PatternFill
        from openpyxl.utils import get_column_letter
        c = self.wb[aba].cell(linha, col)
        c.fill = PatternFill("solid", fgColor=cor[0])
        c.font = Font(color=cor[1], bold=cor[2])
        nome = {v: k for k, v in CORES.items()}.get(cor, cor)
        print(f"      [teste] {aba}!{get_column_letter(col)}{linha} pintada de {nome}")

    def tirar_cor(self, aba, linha, col):
        from openpyxl.styles import PatternFill
        from openpyxl.utils import get_column_letter
        c = self.wb[aba].cell(linha, col)
        if c.fill is not None and c.fill.fill_type:
            c.fill = PatternFill(fill_type=None)
            print(f"      [teste] {aba}!{get_column_letter(col)}{linha} sem cor (dentro do limite)")

    def salvar(self):
        pass


def conectar_excel(cfg):
    """Acha a planilha ABERTA em qualquer janela do Excel, pelo nome do arquivo."""
    import pythoncom
    import win32com.client as win32
    alvo = re.split(r"[\\/]", cfg["planilha"])[-1].lower()
    apps = []
    try:
        apps.append(win32.GetActiveObject("Excel.Application"))
    except Exception:
        pass
    try:
        rot, ctx = pythoncom.GetRunningObjectTable(), pythoncom.CreateBindCtx(0)
        for mk in rot:
            try:
                if mk.GetDisplayName(ctx, None).lower().endswith((".xlsx", ".xlsm", ".xlsb", ".xls")):
                    obj = rot.GetObject(mk).QueryInterface(pythoncom.IID_IDispatch)
                    apps.append(win32.Dispatch(obj).Application)
            except Exception:
                continue
    except Exception:
        pass
    for app in apps:
        try:
            for wb in app.Workbooks:
                if fnmatch.fnmatch(wb.Name.lower(), alvo):
                    return PlanilhaExcel(wb)
            for pv in app.ProtectedViewWindows:
                if fnmatch.fnmatch(pv.Workbook.Name.lower(), alvo):
                    mostrar(f"A planilha {pv.Workbook.Name} está em MODO PROTEGIDO: clique em 'Habilitar Edição'.",
                            chave="protegida", repetir_apos=60)
        except Exception:
            continue
    return None


# =============================================================================
# Regras da etapa 1
# =============================================================================
COLUNAS_E_A_N = ["TICKET DE PESAGEM", "CÓDIGO LACRE", "TRANSPORTADORA", "MOTORISTA", "PLACA DO CAVALO",
                 "UF CAVALO", "PLACA DO REBOQUE", "UF REBOQUE", "MODELO", "PESO_ENTRADA", "HORA_ENTRADA"]
DO_TURNO = ["CÓDIGO LACRE", "TRANSPORTADORA", "MOTORISTA", "PLACA DO CAVALO",
            "UF CAVALO", "PLACA DO REBOQUE", "UF REBOQUE", "MODELO"]


TURNOS = {1: "MATUTINO", 2: "VESPERTINO"}


def peso_alvo(cfg, modelo):
    """Peso bruto alvo do modelo ('LS 4 EIXOS', 'VANDERLEIA'...). None se não achar."""
    m = norm(modelo)
    if not m:
        return None
    alvos = {norm(k): v for k, v in cfg["alvos"].items()}
    if m in alvos:
        return alvos[m]
    for k, v in alvos.items():
        if k and (k in m or m in k):
            return v
    return None


def hora_excel(momento):
    """Hora no formato do Excel (fração do dia): 07:04 -> 0,2944..."""
    return (momento.hour * 3600 + momento.minute * 60) / 86400.0


def tabela_expedicao(tabs):
    for t in tabs:
        if t.col("TICKET DE PESAGEM") is not None:
            return t
    return None


def tabelas_turno(tabs):
    """As tabelas com a coluna PESO A CARREGAR, de cima para baixo: [MATUTINO, VESPERTINO]."""
    return sorted([t for t in tabs if t.col("PESO A CARREGAR") is not None], key=lambda t: t.linha_cab)


def achar_no_turno(cfg, turnos, t):
    """(nº do turno, tabela, índice da linha) do caminhão; tabela=None se não achar."""
    placa = norm(t["cavalo"])
    n = 1 if t["entrada"].time() < cfg["turno2"] else 2
    n = min(n, len(turnos))
    for k in [n] + [k for k in range(1, len(turnos) + 1) if k != n]:
        tab = turnos[k - 1]
        c = tab.col("PLACA DO CAVALO")
        for i, lin in enumerate(tab.linhas):
            if c is not None and norm(lin[c]) == placa:
                return k, tab, i
    return n, None, None


def lancar_tara(planilha, cfg, t):
    """Lança o ticket. Devolve (resultado, aba, linha, observação).
    resultado: LANÇADO | JÁ LANÇADO | ESPERANDO | IGNORADO | ERRO"""
    aba = nome_aba(t["entrada"])
    if aba not in planilha.abas():
        if t["entrada"].date() == dt.date.today():
            return "ESPERANDO", aba, "", f"a aba {aba} ainda não existe na planilha"
        return "IGNORADO", aba, "", f"a planilha não tem a aba {aba}"
    tabs = planilha.tabelas(aba)
    exped, turnos = tabela_expedicao(tabs), tabelas_turno(tabs)
    if exped is None:
        return "ERRO", aba, "", "não achei a tabela de expedição (coluna TICKET DE PESAGEM)"
    if not turnos:
        return "ERRO", aba, "", "não achei as tabelas de turno (coluna PESO A CARREGAR)"

    # 1) caminhão na tabela do turno (matutino/vespertino)
    n, tab_turno, i_turno = achar_no_turno(cfg, turnos, t)
    c_ticket = exped.col("TICKET DE PESAGEM")
    i_existente = next((i for i, lin in enumerate(exped.linhas) if como_int(lin[c_ticket]) == t["numero"]), None)
    linha_existente = exped.linha_excel(i_existente) if i_existente is not None else None
    if tab_turno is None:
        if linha_existente:
            return "JÁ LANÇADO", aba, linha_existente, ""
        return ("ESPERANDO", aba, "",
                f"caminhão {formatar_placa(t['cavalo'])} ({t['motorista']}) não está nas tabelas de turno "
                f"- inclua na tabela do turno {TURNOS.get(n, n)}")
    lin_turno = tab_turno.linhas[i_turno]
    linha_turno = tab_turno.linha_excel(i_turno)

    # 2) PESO A CARREGAR = peso bruto alvo do modelo - peso de entrada
    c_mod, c_carregar = tab_turno.col("MODELO"), tab_turno.col("PESO A CARREGAR")
    modelo = lin_turno[c_mod] if c_mod is not None else None
    alvo = peso_alvo(cfg, modelo)
    peso_entrada = int(round(t["peso_entrada"]))
    carregar = int(round(alvo - peso_entrada)) if alvo else None
    atual = como_int(lin_turno[c_carregar]) if c_carregar is not None else None
    txt_carregar = (f"CARREGAR {kg(carregar)} kg (alvo {kg(alvo)} - entrada {kg(peso_entrada)})" if carregar
                    else f"PESO A CARREGAR NÃO CALCULADO: modelo '{modelo or 'em branco'}' sem peso alvo")

    if linha_existente:
        # ticket já na expedição: só completa o que estiver em branco (peso a carregar e hora de entrada)
        completou = []
        c_hora = exped.col("HORA_ENTRADA")
        if c_hora is not None and vazio(exped.linhas[i_existente][c_hora]):
            planilha.escrever(aba, linha_existente, exped.col_ini + c_hora, [hora_excel(t["entrada"])])
            completou.append(f"hora de entrada {t['entrada']:%H:%M} (linha {linha_existente})")
        if carregar and c_carregar is not None and not atual:
            planilha.escrever(aba, linha_turno, tab_turno.col_ini + c_carregar, [carregar])
            completou.append(txt_carregar)
        if completou:
            return ("COMPLETADO", aba, linha_existente,
                    f"turno {TURNOS.get(n, n)} | {formatar_placa(t['cavalo'])} | " + " | ".join(completou))
        return "JÁ LANÇADO", aba, linha_existente, ""

    valores = {}
    for campo in DO_TURNO:
        c = tab_turno.col(campo)
        v = lin_turno[c] if c is not None else None
        valores[campo] = None if vazio(v) else (v.strip() if isinstance(v, str) else v)
    # o que estiver em branco na tabela do turno, completa com o ticket
    valores["TRANSPORTADORA"] = valores["TRANSPORTADORA"] or t["transportadora"]
    valores["MOTORISTA"] = valores["MOTORISTA"] or t["motorista"]
    valores["PLACA DO CAVALO"] = valores["PLACA DO CAVALO"] or formatar_placa(t["cavalo"])
    valores["PLACA DO REBOQUE"] = valores["PLACA DO REBOQUE"] or (formatar_placa(t["reboque"]) or None)
    valores["TICKET DE PESAGEM"] = t["numero"]
    valores["PESO_ENTRADA"] = peso_entrada
    valores["HORA_ENTRADA"] = hora_excel(t["entrada"])

    # 3) próxima linha livre da expedição (depois da última preenchida)
    chaves = [exped.col(x) for x in ("TICKET DE PESAGEM", "MOTORISTA", "PLACA DO CAVALO", "PESO_ENTRADA")]
    chaves = [c for c in chaves if c is not None]
    ultima = max([i for i, lin in enumerate(exped.linhas) if any(not vazio(lin[c]) for c in chaves)],
                 default=-1)
    if ultima + 1 < len(exped.linhas):
        linha = exped.linha_excel(ultima + 1)
    else:
        linha = planilha.adicionar_linha(aba, exped.nome)
        if exped.col_ini > 1:   # numeração da coluna B
            ant = como_int(planilha.ler(aba, linha - 1, exped.col_ini - 1))
            if ant is not None:
                planilha.escrever(aba, linha, exped.col_ini - 1, [ant + 1])

    # 4) grava E..N e S (só as colunas que existem; colunas vizinhas numa escrita só)
    por_col = {exped.col_ini + exped.col(c): valores[c] for c in COLUNAS_E_A_N if exped.col(c) is not None}
    cols = sorted(por_col)
    inicio = cols[0]
    for k in range(1, len(cols) + 1):
        if k == len(cols) or cols[k] != cols[k - 1] + 1:
            planilha.escrever(aba, linha, inicio, [por_col[c] for c in cols[cols.index(inicio):k]])
            if k < len(cols):
                inicio = cols[k]

    # 5) grava o PESO A CARREGAR na tabela do turno
    if carregar and c_carregar is not None:
        planilha.escrever(aba, linha_turno, tab_turno.col_ini + c_carregar, [carregar])

    obs = (f"turno {TURNOS.get(n, n)} | {valores['MOTORISTA']} | {valores['PLACA DO CAVALO']} | "
           f"{valores['MODELO'] or 'SEM MODELO'} | lacre {valores['CÓDIGO LACRE'] or 'EM BRANCO'} | "
           f"entrada {t['entrada']:%H:%M} | {txt_carregar}")
    return "LANÇADO", aba, linha, obs


# =============================================================================
# Regras da etapa 2 (ticket completo)
# =============================================================================
#            fundo      fonte     negrito
CORES = {"VERDE": ("92D050", "000000", False),
         "AMARELO": ("FFFF00", "000000", False),
         "VERMELHO": ("FF0000", "FFFFFF", True)}


def bipe(cfg):
    if not cfg.get("bipe"):
        return
    try:
        import winsound
        for _ in range(3):
            winsound.MessageBeep(winsound.MB_ICONEXCLAMATION)
            time.sleep(0.4)
    except Exception:
        pass


def lancar_completo(planilha, cfg, t):
    """Peso de saída (O) e hora de saída (T) na linha do ticket + cor da faixa de aceite.
    Devolve (resultado, aba, linha, observação)."""
    aba = nome_aba(t["entrada"])
    if aba not in planilha.abas():
        if t["entrada"].date() == dt.date.today():
            return "ESPERANDO", aba, "", f"a aba {aba} ainda não existe na planilha"
        return "IGNORADO", aba, "", f"a planilha não tem a aba {aba}"
    exped = tabela_expedicao(planilha.tabelas(aba))
    if exped is None:
        return "ERRO", aba, "", "não achei a tabela de expedição (coluna TICKET DE PESAGEM)"

    def achar_linha(exped):
        c_ticket, c_placa, c_saida = (exped.col("TICKET DE PESAGEM"), exped.col("PLACA DO CAVALO"),
                                      exped.col("PESO_SAÍDA"))
        for i, lin in enumerate(exped.linhas):              # 1º: pelo número do ticket
            if como_int(lin[c_ticket]) == t["numero"]:
                return i
        placa = norm(t["cavalo"])                            # 2º: mesma placa ainda sem peso de saída
        cands = [i for i, lin in enumerate(exped.linhas)
                 if c_placa is not None and norm(lin[c_placa]) == placa
                 and (c_saida is None or vazio(lin[c_saida])) and vazio(lin[c_ticket]) is False]
        return cands[-1] if cands else None

    i = achar_linha(exped)
    criou = ""
    if i is None:
        # não chegou o ticket de tara: lança a linha de entrada agora (etapa 1), usando a tara do ticket completo
        r, _, linha_nova, obs = lancar_tara(planilha, cfg, dict(t, peso_entrada=t["tara"]))
        if r not in ("LANÇADO", "JÁ LANÇADO", "COMPLETADO"):
            return r, aba, linha_nova, obs
        exped = tabela_expedicao(planilha.tabelas(aba))
        i = achar_linha(exped)
        if i is None:
            return "ERRO", aba, "", "não consegui achar a linha do ticket depois de lançar a entrada"
        criou = " (linha de entrada criada a partir do ticket completo)"
    lin, linha = exped.linhas[i], exped.linha_excel(i)

    # faixa de aceite: bruto x peso bruto alvo do modelo
    c_mod = exped.col("MODELO")
    modelo = lin[c_mod] if c_mod is not None else None
    alvo = peso_alvo(cfg, modelo)
    bruto = int(round(t["bruto"]))
    if alvo:
        dif = bruto - alvo
        if dif > cfg["tol_exc"]:
            faixa = "VERMELHO"
            situacao = f"EXCESSO DE {kg(dif)} kg (acima de {kg(cfg['tol_exc'])} kg)"
        elif dif > 0:
            faixa = "AMARELO"
            situacao = f"excesso de {kg(dif)} kg (até {kg(cfg['tol_exc'])} kg)"
        elif -dif > cfg["tol_sub"]:
            faixa = "VERDE"
            situacao = f"SUBCARREGADO: {kg(-dif)} kg abaixo do alvo (tolerância {kg(cfg['tol_sub'])} kg)"
        else:
            faixa = None                     # dentro do limite: sem cor
            situacao = "dentro do limite" + (f" ({kg(-dif)} kg abaixo do alvo)" if dif < 0 else " (no alvo)")
        ref = f"bruto {kg(bruto)} x alvo {kg(alvo)} ({modelo})"
    else:
        faixa, situacao, ref = None, f"SEM CONFERÊNCIA: modelo '{modelo or 'em branco'}' sem peso alvo", f"bruto {kg(bruto)}"

    # o que gravar (só o que estiver diferente)
    novos = {"PESO_SAÍDA": bruto, "HORA_SAÍDA": hora_excel(t["saida"])}
    for campo, valor in (("PESO_ENTRADA", int(round(t["tara"]))), ("HORA_ENTRADA", hora_excel(t["entrada"]))):
        c = exped.col(campo)
        if c is not None and vazio(lin[c]):
            novos[campo] = valor                     # completa entrada se estiver em branco
    mudou = {}
    for campo, valor in novos.items():
        c = exped.col(campo)
        if c is None:
            continue
        atual = lin[c]
        if campo.startswith("HORA"):
            if isinstance(atual, dt.datetime):
                atual = atual.time()
            if isinstance(atual, dt.time):
                atual = (atual.hour * 3600 + atual.minute * 60) / 86400.0
            igual = isinstance(atual, (int, float)) and abs(atual % 1 - valor) < 30 / 86400
        else:
            igual = como_int(atual) == valor
        if not igual:
            mudou[campo] = valor
    # cor do PESO_SAÍDA: pinta fora da faixa; dentro do limite tira a cor (corrige cor antiga errada)
    c_saida = exped.col("PESO_SAÍDA")
    if c_saida is not None and alvo:
        if faixa:
            planilha.pintar(aba, linha, exped.col_ini + c_saida, CORES[faixa])
        else:
            planilha.tirar_cor(aba, linha, exped.col_ini + c_saida)
    if not mudou:
        return "JÁ LANÇADO", aba, linha, ""

    for campo, valor in mudou.items():
        planilha.escrever(aba, linha, exped.col_ini + exped.col(campo), [valor])

    alerta = faixa in ("VERMELHO", "VERDE")
    obs = (f"{formatar_placa(t['cavalo'])} | saída {t['saida']:%H:%M} | {ref} | "
           f"{(faixa + ': ') if faixa else ''}{situacao}{criou}")
    return ("SAÍDA - ALERTA" if alerta else "SAÍDA"), aba, linha, obs


# =============================================================================
# Etapa 3 - nota fiscal (DANFE em PDF)
# =============================================================================
class FaltaPypdf(Exception):
    pass


def texto_pdf(caminho):
    try:
        from pypdf import PdfReader
    except ImportError:
        raise FaltaPypdf()
    leitor = PdfReader(caminho)
    return "\n".join((pg.extract_text() or "") for pg in leitor.pages[:3])


_RE_CHAVE = re.compile(r"(?<!\d)(\d{4}(?:[ .]?\d{4}){10})(?!\d)")
_RE_DATA_HORA = re.compile(r"(\d{2})/(\d{2})/(\d{4})\D{0,6}?(\d{2}):(\d{2}):(\d{2})")
_RE_PROTOCOLO = re.compile(r"(?<!\d)(\d{15})(?!\d)\s*[-–]?\s*(\d{2})/(\d{2})/(\d{4})\D{0,6}?(\d{2}):(\d{2}):(\d{2})")
_RE_PLACA = re.compile(r"(?<![A-Z0-9])([A-Z]{3})[ -]?(\d[A-Z0-9]\d{2})(?![A-Z0-9])")
_RE_PESO = re.compile(r"(?<![\d.,])(\d{1,3}(?:\.\d{3})+(?:,\d+)?|\d+,\d{1,4})(?![\d.,])")


def chave_dv_ok(chave):
    """Dígito verificador da chave de acesso: módulo 11, pesos 2..9 da direita."""
    soma, peso = 0, 2
    for ch in reversed(chave[:43]):
        soma += int(ch) * peso
        peso = peso + 1 if peso < 9 else 2
    dv = 11 - soma % 11
    return (0 if dv >= 10 else dv) == int(chave[43])


def ler_nota(caminho, texto=None):
    """Dados da DANFE: nº, data/hora de emissão, tickets, placas e pesos citados. None se não achar o nº."""
    txt = texto if texto is not None else texto_pdf(caminho)
    up = unicodedata.normalize("NFKD", txt)                       # "º" vira "o", "Ç" vira "C"...
    up = "".join(c for c in up if not unicodedata.combining(c)).upper()
    n = {"numero": None, "emissao": None, "tickets": set(), "placas": set(), "pesos": set(), "texto_up": up}

    # nº da nota: pela chave de acesso (44 dígitos, modelo 55); senão, pelo "Nº 000.018.999"
    for m in _RE_CHAVE.finditer(up):
        chave = re.sub(r"\D", "", m.group(1))
        if len(chave) == 44 and chave[20:22] == "55" and chave_dv_ok(chave):
            n["numero"], n["chave"] = int(chave[25:34]), chave
            break
    if not n["numero"]:
        m = re.search(r"N[O0º°]?\.?\s*:?\s*(\d{3}\.\d{3}\.\d{3})", up)
        if m:
            n["numero"] = int(m.group(1).replace(".", ""))

    # HORA DA NOTA: SOMENTE do campo "PROTOCOLO DE AUTORIZAÇÃO DE USO"
    #   (valor padrão: "135260001234567 - 18/09/2026 07:38:16" = nº do protocolo, data e hora)
    n["protocolo"] = None
    m = _RE_PROTOCOLO.search(up)                        # nº do protocolo (15 dígitos) seguido de data e hora
    if not m and "PROTOCOLO DE AUTORIZ" in up:          # ou a 1ª data/hora logo depois do título do campo
        pos = up.find("PROTOCOLO DE AUTORIZ")
        m2 = _RE_DATA_HORA.search(up[pos:pos + 250])
        if m2:
            d, mes, a, h, mi, se = map(int, m2.groups())
            n["protocolo"] = dt.datetime(a, mes, d, h, mi, se)
    elif m:
        d, mes, a, h, mi, se = map(int, m.groups()[1:])
        n["protocolo"] = dt.datetime(a, mes, d, h, mi, se)
    # dia da nota (para achar a aba): o do protocolo; sem protocolo, a DATA DA EMISSÃO
    n["emissao"] = n["protocolo"]
    if not n["emissao"]:
        md = re.search(r"EMISS\w*\D{0,30}(\d{2})/(\d{2})/(\d{4})", up)
        if md:
            d, mes, a = map(int, md.groups())
            n["emissao"] = dt.datetime(a, mes, d)

    n["tickets"] = {int(x) for x in re.findall(r"TICKET\D{0,25}?(\d{4,7})", up)}
    n["placas"] = {a + b for a, b in _RE_PLACA.findall(up)}
    for x in _RE_PESO.findall(up):
        try:
            n["pesos"].add(round(float(x.replace(".", "").replace(",", ".")), 3))
        except ValueError:
            pass
    return n if n["numero"] and n["emissao"] else None


def hora_excel_s(momento):
    return (momento.hour * 3600 + momento.minute * 60 + momento.second) / 86400.0


def lancar_nota(planilha, cfg, n):
    """Nº da nota (D) e hora de emissão (V) na linha do caminhão. Devolve (resultado, aba, linha, obs)."""
    dia = n["emissao"]
    prot = n["protocolo"]                     # hora da nota = protocolo de autorização (None se não tiver)
    abas = [a for a in (nome_aba(dia), nome_aba(dia - dt.timedelta(days=1))) if a in planilha.abas()]
    if not abas:
        return "IGNORADO", nome_aba(dia), "", "sem aba do dia"
    hora = hora_excel_s(prot) if prot else None
    txt_hora = f"hora {prot:%H:%M:%S} (protocolo)" if prot else "SEM PROTOCOLO DE AUTORIZAÇÃO: hora em branco"

    linhas = []   # (aba, tabela, índice)
    for aba in abas:
        exped = tabela_expedicao(planilha.tabelas(aba))
        if exped is None or exped.col("Nº_NOTA") is None:
            continue
        for i in range(len(exped.linhas)):
            linhas.append((aba, exped, i))
    if not linhas:
        return "ERRO", abas[0], "", "não achei a coluna Nº_NOTA na tabela de expedição"

    def val(ex, i, campo):
        c = ex.col(campo)
        return ex.linhas[i][c] if c is not None else None

    # 1) nota já lançada? (completa a hora se estiver em branco)
    for aba, ex, i in linhas:
        if como_int(val(ex, i, "Nº_NOTA")) == n["numero"]:
            c_v = ex.col("HORA EMISSÃO NF")
            if c_v is not None and vazio(ex.linhas[i][c_v]) and hora is not None:
                planilha.escrever(aba, ex.linha_excel(i), ex.col_ini + c_v, [hora])
                return "COMPLETADO", aba, ex.linha_excel(i), f"NF {n['numero']} | {txt_hora}"
            return "JÁ LANÇADO", aba, ex.linha_excel(i), ""

    livres = [(a, ex, i) for a, ex, i in linhas if vazio(val(ex, i, "Nº_NOTA")) and not vazio(val(ex, i, "TICKET DE PESAGEM"))]

    # 2) pelo nº do ticket escrito na nota
    achou, como = None, ""
    for a, ex, i in livres:
        if como_int(val(ex, i, "TICKET DE PESAGEM")) in n["tickets"]:
            achou, como = (a, ex, i), f"pelo nº do ticket {como_int(val(ex, i, 'TICKET DE PESAGEM'))}"
            break
    # 3) pela placa (linha já com peso de saída), desempate pelo peso líquido
    if not achou:
        cands = [(a, ex, i) for a, ex, i in livres
                 if not vazio(val(ex, i, "PESO_SAÍDA"))
                 and (norm(val(ex, i, "PLACA DO CAVALO")) in n["placas"]
                      or norm(val(ex, i, "PLACA DO REBOQUE")) in n["placas"])]
        def bate_peso(c):
            a, ex, i = c
            liq = (como_int(val(ex, i, "PESO_SAÍDA")) or 0) - (como_int(val(ex, i, "PESO_ENTRADA")) or 0)
            return any(abs(p - liq) < 1 or abs(p * 1000 - liq) < 1 for p in n["pesos"])
        com_peso = [c for c in cands if bate_peso(c)]
        if len(cands) > 1 and com_peso:
            cands = com_peso
        if cands:
            achou = cands[0]
            como = "pela placa e peso" if cands[0] in com_peso else "pela placa"
    if not achou:
        motivo = (f"nenhuma linha sem nota com o ticket {sorted(n['tickets']) or '-'} ou a placa "
                  f"{sorted(formatar_placa(p) for p in n['placas'])[:3] or '-'}")
        return "ESPERANDO", abas[0], "", motivo

    aba, ex, i = achou
    linha = ex.linha_excel(i)
    valores = {"Nº_NOTA": n["numero"], "HORA EMISSÃO NF": hora}
    por_col = {ex.col_ini + ex.col(c): v for c, v in valores.items() if ex.col(c) is not None and v is not None}
    for col, v in por_col.items():
        planilha.escrever(aba, linha, col, [v])
    obs = (f"NF {n['numero']} | ticket {como_int(val(ex, i, 'TICKET DE PESAGEM'))} | "
           f"{val(ex, i, 'PLACA DO CAVALO')} | {txt_hora} | ligada {como}")
    return "LANÇADA", aba, linha, obs


_lidas_notas = {}


def ciclo_notas(planilha, cfg):
    """Etapa 3: lê as notas novas e lança nº e hora. Devolve True se gravou algo."""
    pasta = cfg["pasta_notas"]
    if not pasta:
        return False
    try:
        arquivos = listar_pdfs(pasta)
    except OSError as e:
        mostrar(f"Não consegui abrir a pasta das notas: {pasta} ({e})", chave="pasta-notas")
        return False
    limite = time.time() - 60 * 86400
    fila = []
    for e in arquivos:
        mt = e.stat().st_mtime
        if mt < limite:
            continue
        chave = f"{e.path}|{int(mt)}"
        if chave in _resolvidos:
            continue
        if chave not in _lidas_notas:
            try:
                _lidas_notas[chave] = ler_nota(e.path)
            except FaltaPypdf:
                mostrar("Para ler as notas falta a biblioteca pypdf. No Prompt de Comando rode:  pip install pypdf",
                        chave="pypdf", repetir_apos=600)
                return False
            except OSError:
                continue
            except Exception as exc:          # PDF corrompido/protegido
                _lidas_notas[chave] = None
                registrar_erro(exc)
        n = _lidas_notas[chave]
        if n is None:
            if time.time() - mt < 90:
                _lidas_notas.pop(chave)
            elif chave not in _esperando:
                _esperando[chave] = "ilegível"
                mostrar(f"NOTA     NÃO CONSEGUI LER {e.name} (use --ver-nota para conferir)")
                gravar_log("NOTA ERRO", e.name, obs="não achei nº/data da nota no PDF")
            continue
        fila.append((n["emissao"], e, chave, n))
    fila.sort(key=lambda x: (x[0], x[1].name))

    gravou = False
    for _, e, chave, n in fila:
        do_produto = not cfg["produto"] or norm(cfg["produto"]) in norm(n["texto_up"])
        resultado, aba, linha, obs = lancar_nota(planilha, cfg, n)
        if resultado in ("LANÇADA", "COMPLETADO"):
            gravou = True
            _resolvidos.add(chave)
            _esperando.pop(chave, None)
            mostrar(f"NOTA     {resultado:<14} -> aba {aba} linha {linha} | {obs}")
            gravar_log(f"NOTA {resultado}", e.name, n["numero"], aba, linha, obs)
        elif resultado == "JÁ LANÇADO":
            _resolvidos.add(chave)
        elif resultado == "IGNORADO":
            continue
        elif do_produto and _esperando.get(chave) != obs:   # só avisa das notas do produto (concentrado)
            _esperando[chave] = obs
            mostrar(f"NOTA     {resultado:<14} NF {n['numero']} ({e.name}) | {obs}")
            gravar_log(f"NOTA {resultado}", e.name, n["numero"], aba, linha, obs)
    return gravou


def ver_nota(caminho, mostrar_texto=False):
    """--ver-nota: mostra o que o programa leu de uma nota (para conferir sem mandar a nota para ninguém)."""
    try:
        txt = texto_pdf(caminho)
    except FaltaPypdf:
        print("Falta a biblioteca pypdf:  pip install pypdf")
        return
    n = ler_nota(caminho, txt)
    print(f"\nArquivo: {caminho}")
    print(f"Texto extraído: {len(txt)} caracteres")
    if n is None:
        print("NÃO achei o nº e/ou a data/hora da nota.")
        n = ler_nota(caminho, txt + "\n") or {}
    else:
        print(f"Nº da nota ........: {n['numero']}" + (f"   (chave {n['chave']})" if n.get('chave') else ""))
        print(f"Protocolo (hora) ..: " + (f"{n['protocolo']:%d/%m/%Y %H:%M:%S}" if n["protocolo"]
                                             else "NÃO ENCONTRADO -> a hora (coluna V) fica em branco"))
        print(f"Tickets citados ...: {sorted(n['tickets']) or 'nenhum'}")
        print(f"Placas encontradas : {sorted(formatar_placa(p) for p in n['placas']) or 'nenhuma'}")
        print(f"Pesos encontrados .: {sorted(n['pesos'])[:12]}")
    if mostrar_texto:
        print("\n----- TEXTO EXTRAÍDO -----\n" + txt)


# =============================================================================
# Ciclo
# =============================================================================
_lidos = {}        # caminho|versão -> ticket (cada PDF é lido uma vez só)
_resolvidos = set()  # PDFs já lançados/já existentes nesta sessão
_esperando = {}    # PDF -> motivo
_resumo = [None]


def listar_pdfs(pasta, nivel=0):
    res = []
    for e in os.scandir(pasta):
        if e.is_file() and e.name.lower().endswith(".pdf"):
            res.append(e)
        elif e.is_dir() and nivel < 2:
            try:
                res.extend(listar_pdfs(e.path, nivel + 1))
            except OSError:
                pass
    return res


def ciclo(planilha, cfg):
    planilha.limpar_cache()
    pastas = [("Tara", cfg["pasta"]), ("Completos", cfg["pasta_completos"])]
    limite = time.time() - 60 * 86400
    fila, contagem = [], {}
    for rotulo, pasta in pastas:
        if not pasta:
            continue
        try:
            arquivos = listar_pdfs(pasta)
        except OSError as e:
            mostrar(f"Não consegui abrir a pasta {rotulo}: {pasta} ({e})", chave=f"pasta-{rotulo}")
            continue
        n_total = 0
        for e in arquivos:
            mt = e.stat().st_mtime
            if mt < limite:
                continue
            n_total += 1
            chave = f"{e.path}|{int(mt)}"
            if chave in _resolvidos:
                continue
            if chave not in _lidos:
                try:
                    _lidos[chave] = ler_ticket(e.path)
                except OSError:
                    continue            # arquivo sendo gravado ou bloqueado: tenta no próximo ciclo
                except Exception as exc:
                    # PDF com conteudo inesperado: conta como ilegivel. Deixar o erro
                    # subir abortaria o ciclo inteiro, a cada intervalo, para sempre.
                    _lidos[chave] = None
                    registrar_erro(exc)
            t = _lidos[chave]
            if t is None:
                if time.time() - mt < 90:
                    _lidos.pop(chave)   # pode estar sendo gravado: lê de novo depois
                elif chave not in _esperando:
                    _esperando[chave] = "ilegível"
                    mostrar(f"NÃO CONSEGUI LER {e.name} (layout diferente do esperado?)")
                    gravar_log("ERRO", e.name, obs="não consegui ler os dados do PDF")
                continue
            if cfg["produto"] and norm(cfg["produto"]) not in norm(t["produto"]):
                _resolvidos.add(chave)
                continue
            momento = t["saida"] if t["tipo"] == "completo" else t["entrada"]
            fila.append((momento, 0 if t["tipo"] == "tara" else 1, e, chave, t))
        contagem[rotulo] = n_total
    fila.sort(key=lambda x: (x[0], x[1], x[2].name))   # na ordem em que as pesagens aconteceram

    por_resultado = {}
    for _, _, e, chave, t in fila:
        if t["tipo"] == "tara":
            resultado, aba, linha, obs = lancar_tara(planilha, cfg, t)
        else:
            resultado, aba, linha, obs = lancar_completo(planilha, cfg, t)
        etiqueta = "TARA" if t["tipo"] == "tara" else "COMPLETO"
        por_resultado[resultado] = por_resultado.get(resultado, 0) + 1
        if resultado in ("LANÇADO", "COMPLETADO", "SAÍDA", "SAÍDA - ALERTA"):
            _resolvidos.add(chave)
            _esperando.pop(chave, None)
            mostrar(f"{etiqueta:<8} {resultado:<14} ticket {t['numero']} -> aba {aba} linha {linha} | {obs}")
            gravar_log(f"{etiqueta} {resultado}", e.name, t["numero"], aba, linha, obs)
            if resultado == "SAÍDA - ALERTA":
                bipe(cfg)
        elif resultado in ("JÁ LANÇADO", "IGNORADO"):
            _resolvidos.add(chave) if resultado == "JÁ LANÇADO" else None
        else:  # ESPERANDO / ERRO: avisa uma vez e tenta de novo nos próximos ciclos
            if _esperando.get(chave) != obs:
                _esperando[chave] = obs
                mostrar(f"{etiqueta:<8} {resultado:<14} ticket {t['numero']} ({e.name}) | {obs}")
                gravar_log(f"{etiqueta} {resultado}", e.name, t["numero"], aba, linha, obs)

    resumo = (tuple(sorted(contagem.items())), len([k for k in _esperando if k not in _resolvidos]),
              por_resultado.get("IGNORADO", 0))
    if resumo != _resumo[0]:
        _resumo[0] = resumo
        mostrar(" | ".join(f"Pasta {r}: {n} PDF(s)" for r, n in sorted(contagem.items()))
                + f" | {resumo[1]} esperando | {resumo[2]} de dias sem aba nesta planilha")
    ciclo_notas(planilha, cfg)
    # salva pelo que ficou PENDENTE, nao pelo que este ciclo gravou: se um ciclo
    # anterior caiu com erro depois de escrever, o que ele escreveu salva agora
    if cfg["salvar"] and getattr(planilha, "pendente", False):
        planilha.salvar()


# =============================================================================
# Erros do Excel
# =============================================================================
_OCUPADO = {-2147418111, -2147417846, -2146777998}   # célula em edição / janela aberta no Excel


def excel_ocupado(exc):
    if type(exc).__name__ != "com_error":
        return False
    codigos = {exc.args[0]}
    if len(exc.args) > 2 and exc.args[2]:
        codigos.add(exc.args[2][5])
    return bool(codigos & _OCUPADO)


def descrever(exc):
    if type(exc).__name__ == "com_error":
        info = exc.args[2] if len(exc.args) > 2 else None
        texto = f"COM {exc.args[0]}: {exc.args[1]}" + (f" | {info[2]}" if info and info[2] else "")
    else:
        texto = f"{type(exc).__name__}: {exc}"
    quadros = [q for q in traceback.extract_tb(exc.__traceback__) if q.filename == os.path.abspath(__file__)]
    return texto + (f" [em {quadros[-1].name}, linha {quadros[-1].lineno}]" if quadros else "")


def registrar_erro(exc):
    try:
        with open(ARQ_ERROS, "a", encoding="utf-8") as f:
            f.write(f"\n===== {dt.datetime.now():%d/%m/%Y %H:%M:%S} =====\n{descrever(exc)}\n")
            f.write("".join(traceback.format_exception(type(exc), exc, exc.__traceback__)))
    except OSError:
        pass


# =============================================================================
# Programa
# =============================================================================
def main():
    ap = argparse.ArgumentParser(description="Etapas 1, 2 e 3: tickets de tara, completos e notas fiscais.")
    ap.add_argument("--teste", metavar="ARQUIVO_XLSX", help="só mostra o que faria, sem gravar nada")
    ap.add_argument("--ver-nota", metavar="NOTA_PDF", help="mostra o que o programa leu de uma nota fiscal")
    ap.add_argument("--ver-nota-texto", metavar="NOTA_PDF", help="igual, e mostra também o texto extraído")
    args = ap.parse_args()
    if args.ver_nota or args.ver_nota_texto:
        ver_nota(args.ver_nota or args.ver_nota_texto, mostrar_texto=bool(args.ver_nota_texto))
        return
    cfg = carregar_config()
    print("=" * 78)
    print(" LANÇADOR - ETAPA 1 (tara) + ETAPA 2 (completo) + ETAPA 3 (nota fiscal: D e V)")
    print(f" Pasta dos tickets de tara:      {cfg['pasta']}")
    print(f" Pasta dos tickets completos:    {cfg['pasta_completos']}")
    print(f" Pasta das notas fiscais:        {cfg['pasta_notas']}")
    print(f" Planilha:                  {cfg['planilha']}")
    print(f" Matutino até {cfg['turno2']:%H:%M}, vespertino depois | Alvos: "
          + ", ".join(f"{k} {kg(v)}" for k, v in cfg["alvos"].items()))
    print(f" Tolerância: +{kg(cfg['tol_exc'])} kg / -{kg(cfg['tol_sub'])} kg | "
          f"VERDE subcarga > {kg(cfg['tol_sub'])} kg, AMARELO excesso até {kg(cfg['tol_exc'])} kg, VERMELHO acima")
    print(f" Intervalo: {cfg['intervalo']} s  (Ctrl+C para parar)")
    print("=" * 78, flush=True)

    if args.teste:
        planilha = PlanilhaTeste(args.teste)
        mostrar(f"MODO TESTE em {planilha.nome}: nada será gravado")
        ciclo(planilha, cfg)
        return

    try:
        import win32com.client  # noqa: F401
    except ImportError:
        print("\nFalta a biblioteca pywin32. No Prompt de Comando rode:  pip install pywin32\n")
        if sys.stdin is not None and sys.stdin.isatty():
            input("Enter para sair...")
        return

    planilha = None
    while True:
        try:
            if planilha is None or not planilha.viva():
                planilha = conectar_excel(cfg)
                if planilha is None:
                    mostrar("Aguardando a planilha ser aberta no Excel...", chave="aguardando", repetir_apos=120)
                else:
                    abas = [a for a in planilha.abas() if re.fullmatch(r"\d{2}\.\d{2}", a)]
                    mostrar(f"Conectado à planilha {planilha.nome} | abas de dia: "
                            f"{abas[0] + ' até ' + abas[-1] if abas else 'nenhuma'}")
            if planilha is not None:
                ciclo(planilha, cfg)
        except KeyboardInterrupt:
            break
        except Exception as exc:
            if excel_ocupado(exc):
                mostrar("Excel ocupado (célula em edição ou janela aberta). Tento de novo em instantes.",
                        chave="ocupado", repetir_apos=60)
            else:
                mostrar(f"ERRO: {descrever(exc)}  (detalhes em etapa3_erros.log)", chave=str(exc), repetir_apos=120)
                registrar_erro(exc)
            if planilha is not None and not planilha.viva():
                planilha = None
        try:
            time.sleep(cfg["intervalo"])
        except KeyboardInterrupt:
            break
    mostrar("Programa encerrado.")


if __name__ == "__main__":
    main()
